# -*- coding: utf-8 -*-
"""Veículos da unidade: cadastro próprio (fora dos equipamentos), usos do mês e RDV."""
import io
from datetime import date, datetime

from flask import Blueprint, abort, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app import db
from app.models.agenda import AgendaEvento
from app.models.unidade import Unidade
from app.models.veiculo import (CATEGORIAS_VEICULO, COMBUSTIVEIS_VEICULO, PREFIXO_MAX_DIGITOS, ModeloVeiculo,
                                Veiculo, normalizar_placa, normalizar_prefixo)
from app.services import reservas
from app.utils import agora_brasilia, agora_local

veiculos_bp = Blueprint('veiculos', __name__, url_prefix='/veiculos')

MESES = ['Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho', 'Julho',
         'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro']


def _acesso_unidade(unidade_id):
    if current_user.pode('ver_todas_unidades'):
        return True
    return unidade_id in {uu.unidade_id for uu in current_user.unidades.filter_by(ativo=True).all()}


def _veiculo_ou_404(id):
    veiculo = Veiculo.query.get_or_404(id)
    if not _acesso_unidade(veiculo.unidade_id):
        abort(403)
    return veiculo


def _mes_param():
    texto = request.values.get('mes') or ''
    try:
        ano, mes = (int(p) for p in texto.split('-')[:2])
        return date(ano, mes, 1)
    except (TypeError, ValueError):
        hoje = agora_brasilia().date()
        return date(hoje.year, hoje.month, 1)


def _proximo_mes(d):
    return date(d.year + (d.month == 12), d.month % 12 + 1, 1)


def _usos_do_mes(veiculo, mes):
    ini = datetime.combine(mes, datetime.min.time())
    fim = datetime.combine(_proximo_mes(mes), datetime.min.time())
    return (AgendaEvento.query
            .filter(AgendaEvento.veiculo_id == veiculo.id, AgendaEvento.inicio >= ini, AgendaEvento.inicio < fim)
            .order_by(AgendaEvento.inicio)
            .all())


def _finalidade(ev, rotulo):
    titulo = ev.titulo or ''
    return '' if titulo.startswith(rotulo) else titulo


def _inteiro(texto):
    digitos = ''.join(ch for ch in (texto or '') if ch.isdigit())
    return int(digitos) if digitos else None


def _data(texto):
    try:
        return datetime.strptime((texto or '').strip(), '%Y-%m-%d').date()
    except ValueError:
        return None


def _duplicado(veiculo, placa, prefixo, alugado):
    """Placa e prefixo não se repetem entre veículos ativos (383 e AL-383 são carros diferentes)."""
    outros = Veiculo.query.filter(Veiculo.ativo.is_(True), Veiculo.id != (veiculo.id or 0))
    mesmo = outros.filter(Veiculo.placa == placa).first()
    if mesmo:
        return f'Já existe um veículo ativo com a placa {mesmo.placa_exibicao} ({mesmo.unidade.nome}).'
    mesmo = outros.filter(Veiculo.prefixo == prefixo, Veiculo.alugado.is_(alugado)).first()
    if mesmo:
        return f'O prefixo {mesmo.prefixo_exibicao} já é do veículo {mesmo.placa_exibicao} ({mesmo.unidade.nome}).'
    return None


def _preencher(veiculo, form):
    """Copia o formulário para o veículo. Devolve a mensagem de erro, se houver."""
    prefixo, digitou_al = normalizar_prefixo(form.get('prefixo'))
    placa = normalizar_placa(form.get('placa'))
    if not prefixo:
        return 'Informe o prefixo do veículo.'
    if not prefixo.isdigit() or len(prefixo) > PREFIXO_MAX_DIGITOS:
        return f'O prefixo tem só números, até {PREFIXO_MAX_DIGITOS} dígitos (ex.: 383 ou 1024). Alugado é o check abaixo.'
    if len(placa) != 7:
        return 'Informe a placa com 7 caracteres (ex.: ABC1D23 ou ABC1234).'
    categoria = form.get('categoria') or ''
    if categoria not in CATEGORIAS_VEICULO:
        return 'Escolha a categoria do veículo.'
    alugado = digitou_al or form.get('alugado') == 'on'
    erro = _duplicado(veiculo, placa, prefixo, alugado)
    if erro:
        return erro

    combustivel = form.get('combustivel') or ''
    veiculo.prefixo = prefixo
    veiculo.alugado = alugado
    veiculo.locadora = ((form.get('locadora') or '').strip()[:150] or None) if alugado else None
    veiculo.placa = placa
    veiculo.marca = (form.get('marca') or '').strip()[:60] or None
    veiculo.modelo = (form.get('modelo') or '').strip()[:80] or None
    veiculo.categoria = categoria
    veiculo.ano = (form.get('ano') or '').strip()[:9] or None
    veiculo.cor = (form.get('cor') or '').strip()[:30] or None
    veiculo.combustivel = combustivel if combustivel in COMBUSTIVEIS_VEICULO else None
    veiculo.lotacao = _inteiro(form.get('lotacao'))
    veiculo.renavam = ''.join(ch for ch in (form.get('renavam') or '') if ch.isdigit())[:20] or None
    veiculo.chassi = normalizar_placa(form.get('chassi'))[:30] or None
    veiculo.km_cadastro = _inteiro(form.get('km_cadastro'))
    veiculo.licenciamento_vencimento = _data(form.get('licenciamento_vencimento'))
    veiculo.observacoes = (form.get('observacoes') or '').strip() or None
    return None


def _render_form(unidade, veiculo):
    modelos = ModeloVeiculo.query.order_by(ModeloVeiculo.marca, ModeloVeiculo.nome).all()
    return render_template(
        'veiculos/form.html', unidade=unidade, veiculo=veiculo, editando=veiculo.id is not None,
        dados_form=request.form if request.method == 'POST' else None,
        categorias=CATEGORIAS_VEICULO, combustiveis=COMBUSTIVEIS_VEICULO,
        marcas=sorted({m.marca for m in modelos}),
        modelos=[{'marca': m.marca, 'nome': m.nome, 'categoria': m.categoria} for m in modelos],
    )


@veiculos_bp.route('/novo/<int:unidade_id>', methods=['GET', 'POST'])
@login_required
def novo(unidade_id):
    if not current_user.pode('cadastrar_veiculo'):
        abort(403)
    unidade = Unidade.query.get_or_404(unidade_id)
    if not _acesso_unidade(unidade.id):
        abort(403)
    veiculo = Veiculo(unidade_id=unidade.id, criado_por=current_user.id)
    if request.method == 'POST':
        erro = _preencher(veiculo, request.form)
        if erro:
            flash(erro, 'danger')
        else:
            db.session.add(veiculo)
            db.session.commit()
            flash(f'Veículo {veiculo.prefixo_exibicao} cadastrado. Ele já pode ser reservado na agenda.', 'success')
            return redirect(url_for('unidades.detalhe', id=unidade.id, _anchor='tab-veiculos'))
    return _render_form(unidade, veiculo)


@veiculos_bp.route('/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def editar(id):
    if not current_user.pode('editar_veiculo'):
        abort(403)
    veiculo = _veiculo_ou_404(id)
    if request.method == 'POST':
        erro = _preencher(veiculo, request.form)
        if erro:
            flash(erro, 'danger')
        else:
            db.session.commit()
            flash('Veículo atualizado.', 'success')
            return redirect(url_for('veiculos.usos', id=veiculo.id))
    return _render_form(veiculo.unidade, veiculo)


@veiculos_bp.route('/<int:id>/ativo', methods=['POST'])
@login_required
def alternar_ativo(id):
    if not current_user.pode('editar_veiculo'):
        abort(403)
    veiculo = _veiculo_ou_404(id)
    if not veiculo.ativo:
        erro = _duplicado(veiculo, veiculo.placa, veiculo.prefixo, veiculo.alugado)
        if erro:
            flash(f'Não deu para reativar: {erro}', 'danger')
            return redirect(url_for('veiculos.usos', id=veiculo.id))
    veiculo.ativo = not veiculo.ativo
    db.session.commit()
    flash('Veículo reativado.' if veiculo.ativo else
          'Veículo desativado: sai da lista da unidade e da agenda, mas o histórico de usos continua aqui.', 'success')
    return redirect(url_for('veiculos.usos', id=veiculo.id))


@veiculos_bp.route('/<int:id>')
@login_required
def usos(id):
    veiculo = _veiculo_ou_404(id)
    mes = _mes_param()
    rotulo = veiculo.rotulo
    lista = _usos_do_mes(veiculo, mes)
    agora = agora_brasilia()
    situacao = reservas.situacao_veiculos([veiculo], agora)[veiculo.id]
    anterior = date(mes.year - (mes.month == 1), (mes.month - 2) % 12 + 1, 1)
    return render_template(
        'veiculos/usos.html',
        veiculo=veiculo, rotulo=rotulo, usos=lista, mes=mes,
        mes_nome=f'{MESES[mes.month - 1]}/{mes.year}',
        mes_anterior=anterior.strftime('%Y-%m'), mes_seguinte=_proximo_mes(mes).strftime('%Y-%m'),
        situacao=situacao, km_atual=reservas.km_atual(veiculo),
        total_km=sum(ev.km_rodados or 0 for ev in lista),
        finalidade=lambda ev: _finalidade(ev, rotulo),
        agora=agora,
        pode_reservar=veiculo.ativo and current_user.pode('adicionar_agenda'),
        pode_editar=current_user.pode('editar_veiculo'),
    )


@veiculos_bp.route('/<int:id>/km/<int:evento_id>', methods=['POST'])
@login_required
def salvar_km(id, evento_id):
    veiculo = _veiculo_ou_404(id)
    ev = AgendaEvento.query.filter_by(id=evento_id, veiculo_id=veiculo.id).first_or_404()

    def _km(nome):
        bruto = ''.join(ch for ch in (request.form.get(nome) or '') if ch.isdigit())
        return int(bruto) if bruto else None

    km_saida, km_chegada = _km('km_saida'), _km('km_chegada')
    if km_saida is not None and km_chegada is not None and km_chegada < km_saida:
        flash('O km de chegada não pode ser menor que o de saída.', 'danger')
    else:
        ev.km_saida = km_saida
        ev.km_chegada = km_chegada
        ev.devolvido_em = agora_local() if km_chegada is not None else None
        ev.atualizado_em = agora_local()
        db.session.commit()
        flash('Km registrado.', 'success')
    return redirect(url_for('veiculos.usos', id=veiculo.id, mes=ev.inicio.strftime('%Y-%m')))


@veiculos_bp.route('/<int:id>/rdv')
@login_required
def rdv(id):
    """RDV do mês em Excel, já preenchido com as reservas da agenda."""
    veiculo = _veiculo_ou_404(id)
    mes = _mes_param()
    rotulo = veiculo.rotulo
    lista = _usos_do_mes(veiculo, mes)
    unidade = veiculo.unidade

    wb = Workbook()
    ws = wb.active
    ws.title = f'RDV {mes.strftime("%m-%Y")}'
    escuro = PatternFill('solid', fgColor='0D3B5E')
    claro = PatternFill('solid', fgColor='E8F2F8')
    fino = Side(style='thin', color='A0AEC0')
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    negrito = Font(bold=True, color='1E3A50')
    colunas = [('Data', 11), ('Condutor', 26), ('Destino / itinerário', 30), ('Finalidade', 26),
               ('Hora saída', 10), ('Km saída', 11), ('Hora chegada', 11), ('Km chegada', 11),
               ('Km rodados', 11), ('Assinatura do condutor', 26)]
    ultima = get_column_letter(len(colunas))
    for i, (_, largura) in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura

    def _faixa(linha, texto, fonte, preenchimento=None, altura=None):
        ws.merge_cells(f'A{linha}:{ultima}{linha}')
        c = ws[f'A{linha}']
        c.value = texto
        c.font = fonte
        c.alignment = Alignment(horizontal='center', vertical='center')
        if preenchimento:
            c.fill = preenchimento
        if altura:
            ws.row_dimensions[linha].height = altura

    _faixa(1, 'PREFEITURA DE SOROCABA · SECRETARIA DA SAÚDE', Font(bold=True, size=11, color='FFFFFF'), escuro, 20)
    _faixa(2, 'RDV · RELATÓRIO DIÁRIO DE VEÍCULO', Font(bold=True, size=14, color='0D3B5E'), None, 24)

    dados = [
        ('Unidade', unidade.nome, 'Mês/ano', f'{MESES[mes.month - 1]}/{mes.year}'),
        ('Veículo', veiculo.marca_modelo or veiculo.categoria or '', 'Placa', veiculo.placa_exibicao),
        ('Prefixo', veiculo.prefixo_exibicao, 'Ano', veiculo.ano or ''),
        ('Categoria', veiculo.categoria or '', 'Combustível', veiculo.combustivel or ''),
    ]
    linha = 4
    for r1, v1, r2, v2 in dados:
        ws.cell(row=linha, column=1, value=r1).font = negrito
        ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=4)
        ws.cell(row=linha, column=2, value=v1)
        ws.cell(row=linha, column=5, value=r2).font = negrito
        ws.merge_cells(start_row=linha, start_column=6, end_row=linha, end_column=len(colunas))
        ws.cell(row=linha, column=6, value=v2)
        linha += 1

    linha += 1
    cab = linha
    for i, (nome, _) in enumerate(colunas, start=1):
        c = ws.cell(row=cab, column=i, value=nome)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1A82B8')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = borda
    ws.row_dimensions[cab].height = 30

    linhas_uso = max(len(lista), 31)
    for n in range(linhas_uso):
        r = cab + 1 + n
        ev = lista[n] if n < len(lista) else None
        if ev:
            ini, fim = reservas.intervalo_evento(ev)
            valores = [
                ev.inicio.date(),
                ev.condutor.nome if ev.condutor else (ev.criador.nome if ev.criador else ''),
                ev.local or '',
                _finalidade(ev, rotulo),
                None if ev.dia_inteiro else ini.strftime('%H:%M'),
                ev.km_saida,
                None if ev.dia_inteiro else fim.strftime('%H:%M'),
                ev.km_chegada,
            ]
        else:
            valores = [None] * 8
        for i, v in enumerate(valores, start=1):
            ws.cell(row=r, column=i, value=v)
        ws.cell(row=r, column=9, value=f'=IF(AND(F{r}<>"",H{r}<>""),H{r}-F{r},"")')
        for i in range(1, len(colunas) + 1):
            c = ws.cell(row=r, column=i)
            c.border = borda
            c.alignment = Alignment(vertical='center', horizontal='center' if i not in (2, 3, 4, 10) else 'left',
                                    wrap_text=i in (3, 4))
            if n % 2:
                c.fill = claro
        ws.cell(row=r, column=1).number_format = 'DD/MM/YYYY'
        for col in (6, 8, 9):
            ws.cell(row=r, column=col).number_format = '#,##0'
        ws.row_dimensions[r].height = 20

    total = cab + linhas_uso + 1
    ws.cell(row=total, column=8, value='Total km').font = negrito
    c = ws.cell(row=total, column=9, value=f'=SUM(I{cab + 1}:I{total - 1})')
    c.font = negrito
    c.number_format = '#,##0'
    c.border = borda
    ws.cell(row=total + 3, column=1, value='Responsável pela unidade: ________________________________')
    ws.cell(row=total + 3, column=6, value='Data: ____/____/________')
    ws.cell(row=total + 5, column=1,
            value=f'Gerado pelo SIGUS em {agora_brasilia().strftime("%d/%m/%Y %H:%M")} a partir das reservas da agenda.'
            ).font = Font(italic=True, size=8, color='718096')

    ws.freeze_panes = ws.cell(row=cab + 1, column=1)
    ws.print_title_rows = f'{cab}:{cab}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True,
        download_name=f'RDV_{veiculo.prefixo_exibicao}_{veiculo.placa}_{mes.strftime("%Y-%m")}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
