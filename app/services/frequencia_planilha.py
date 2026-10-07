# -*- coding: utf-8 -*-
"""Leitura da planilha de frequência mensal enviada ao RH (.xls ou .xlsx).

Tudo é ligado pela matrícula: nas abas fora da Capa o nome e a função são fórmulas
que o arquivo nem sempre traz calculadas.
"""
import calendar
import io
import re
import unicodedata
from datetime import date, datetime

from app.models.frequencia import MESES_PT, normalizar_matricula


class PlanilhaInvalida(Exception):
    pass


def _norm(texto):
    if texto is None:
        return ''
    texto = unicodedata.normalize('NFKD', str(texto))
    texto = ''.join(ch for ch in texto if not unicodedata.combining(ch))
    texto = texto.replace('.', ' ').upper()
    return re.sub(r'\s+', ' ', texto).strip()


def _ler_xls(conteudo):
    import xlrd
    wb = xlrd.open_workbook(file_contents=conteudo)
    abas = {}
    for sh in wb.sheets():
        linhas = []
        for r in range(sh.nrows):
            linha = []
            for c in range(sh.ncols):
                cel = sh.cell(r, c)
                if cel.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    linha.append(None)
                elif cel.ctype == xlrd.XL_CELL_TEXT:
                    linha.append(cel.value.strip() or None)
                elif cel.ctype == xlrd.XL_CELL_DATE:
                    try:
                        linha.append(xlrd.xldate_as_datetime(cel.value, wb.datemode))
                    except Exception:
                        linha.append(None)
                else:
                    linha.append(cel.value)
            linhas.append(linha)
        abas[_norm(sh.name)] = linhas
    return abas


def _ler_xlsx(conteudo):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    abas = {}
    for ws in wb.worksheets:
        linhas = []
        for row in ws.iter_rows(values_only=True):
            linhas.append([
                (v.strip() or None) if isinstance(v, str) else v
                for v in row
            ])
        abas[_norm(ws.title)] = linhas
    wb.close()
    return abas


def _cel(linha, col):
    if col is None or col >= len(linha):
        return None
    return linha[col]


def _num(valor):
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    txt = str(valor).strip().replace(',', '.')
    try:
        return float(txt)
    except ValueError:
        return None


_LIMITES = {'matricula': 20, 'nome': 200, 'funcao': 120, 'local': 120, 'regime': 80,
            'condicao': 40, 'tipo_jornada': 20}


def _txt(valor):
    """Texto de célula; números soltos (fórmula sem cálculo = 0) viram None."""
    if valor is None or isinstance(valor, (int, float)):
        return None
    txt = re.sub(r'\s+', ' ', str(valor)).strip()
    return txt or None


def _achar_cabecalho(linhas, exigidos, limite=40):
    """Primeira linha que contém todos os rótulos exigidos. Devolve (índice, {rótulo: coluna})."""
    for i, linha in enumerate(linhas[:limite]):
        mapa = {}
        for c, v in enumerate(linha):
            if isinstance(v, str):
                mapa.setdefault(_norm(v), c)
        if all(any(chave.startswith(e) for chave in mapa) for e in exigidos):
            return i, mapa
    return None, {}


def _col(mapa, *nomes):
    for nome in nomes:
        if nome in mapa:
            return mapa[nome]
    for nome in nomes:
        for chave, c in mapa.items():
            if chave.startswith(nome):
                return c
    return None


def _aba(abas, *nomes):
    for nome in nomes:
        if nome in abas:
            return abas[nome]
    return None


_RE_COMPETENCIA = re.compile(r'(' + '|'.join(_norm(m) for m in MESES_PT) + r')\s*/\s*(\d{4})')


def _competencia(linhas):
    for linha in linhas:
        for v in linha:
            if isinstance(v, str):
                m = _RE_COMPETENCIA.search(_norm(v))
                if m:
                    mes = [_norm(x) for x in MESES_PT].index(m.group(1)) + 1
                    return date(int(m.group(2)), mes, 1)
            elif isinstance(v, datetime) and v.year > 1990:
                return date(v.year, v.month, 1)
    return None


def _local(linhas, nome_arquivo):
    candidatos = []
    for linha in linhas:
        for v in linha:
            if isinstance(v, str):
                n = _norm(v)
                if n.startswith('FREQUENCIA') or _RE_COMPETENCIA.search(n):
                    continue
                candidatos.append(re.sub(r'\s+', ' ', v).strip())
    for c in candidatos:
        if _norm(c).startswith('SES'):
            return c.upper()
    if candidatos:
        return max(candidatos, key=len).upper()
    base = re.sub(r'\.(xlsx?|xlsm)$', '', nome_arquivo or '', flags=re.I).strip()
    return base.upper() or None


_TOTAIS_CAPA = {
    'FA': ('FALTA ABON',),
    'AM': ('ATEST MED',),
    'LTPF': ('LTPF',),
    'JE': ('JUST ELEIT',),
    'DS': ('DOAC SANG',),
    'LN': ('LICEN NOJO',),
    'LG': ('LICEN GALA',),
    'LP': ('LICEN PATER',),
    'FJ': ('FALTA JUST',),
    'FI': ('FALTA INJUST',),
    'DSR': ('FALTA DSR',),
}


def ler_planilha(conteudo, nome_arquivo):
    """Converte a planilha em dados. Levanta PlanilhaInvalida com mensagem amigável."""
    ext = (nome_arquivo or '').lower().rsplit('.', 1)[-1]
    try:
        if ext == 'xls':
            abas = _ler_xls(conteudo)
        elif ext in ('xlsx', 'xlsm'):
            abas = _ler_xlsx(conteudo)
        else:
            raise PlanilhaInvalida('Envie a planilha no formato .xls ou .xlsx.')
    except PlanilhaInvalida:
        raise
    except Exception:
        raise PlanilhaInvalida('Não consegui abrir o arquivo. Confira se é a planilha de frequência do RH.')

    capa = _aba(abas, 'CAPA')
    if not capa:
        raise PlanilhaInvalida('Não achei a aba "Capa". Essa planilha não parece ser a frequência mensal do RH.')

    i_cab, cab = _achar_cabecalho(capa, ['MATR', 'NOME'])
    if i_cab is None:
        raise PlanilhaInvalida('Não achei o cabeçalho (MATR / NOME) na aba "Capa".')

    competencia = _competencia(capa[:i_cab]) or _competencia(capa)
    if not competencia:
        raise PlanilhaInvalida('Não achei o mês de referência (ex.: JULHO/2026) na aba "Capa".')
    local = _local(capa[:i_cab], nome_arquivo)
    if not local:
        raise PlanilhaInvalida('Não achei o nome da unidade na aba "Capa".')

    avisos = []
    lancamentos = {}

    def _lanc(mat):
        if mat not in lancamentos:
            lancamentos[mat] = {
                'matricula': mat, 'nome': None, 'funcao': None, 'regime': None, 'condicao': None,
                'tipo_jornada': None, 'horas_mensais': None, 'vencimento': None, 'hora_noturna': None,
                'totais': {}, 'dias': {}, 'complemento': None, 'horario': None,
                'banco_saldo_anterior': None, 'banco_realizadas': None, 'banco_utilizadas': None,
                'banco_saldo_atual': None, 'banco_obs': None,
            }
        return lancamentos[mat]

    c_mat = _col(cab, 'MATR')
    c_nome = _col(cab, 'NOME')
    c_func = _col(cab, 'FUNCAO')
    c_reg = _col(cab, 'REGIME')
    c_cond = _col(cab, 'COND')
    c_hm = _col(cab, 'HM')
    c_tipo = _col(cab, 'TIPO')
    c_venc = _col(cab, 'VENC')
    c_hnot = _col(cab, 'HORA NOT')
    c_totais = {sigla: _col(cab, *rotulos) for sigla, rotulos in _TOTAIS_CAPA.items()}
    for linha in capa[i_cab + 1:]:
        mat = normalizar_matricula(_cel(linha, c_mat))
        if not mat:
            continue
        l = _lanc(mat)
        l['nome'] = _txt(_cel(linha, c_nome))
        l['funcao'] = _txt(_cel(linha, c_func))
        l['regime'] = _txt(_cel(linha, c_reg))
        l['condicao'] = _txt(_cel(linha, c_cond))
        l['tipo_jornada'] = _txt(_cel(linha, c_tipo))
        l['horas_mensais'] = _num(_cel(linha, c_hm))
        l['vencimento'] = _num(_cel(linha, c_venc))
        l['hora_noturna'] = _num(_cel(linha, c_hnot)) or None
        for sigla, c in c_totais.items():
            v = _num(_cel(linha, c))
            if v:
                l['totais'][sigla] = v

    just = _aba(abas, 'JUSTIFICATIVAS')
    if just:
        i_j, cab_j = _achar_cabecalho(just, ['MATR', 'OCORRENCIAS'])
        if i_j is not None:
            ultimo_dia = calendar.monthrange(competencia.year, competencia.month)[1]
            dias_col = {}
            for c, v in enumerate(just[i_j]):
                n = _num(v)
                if n is not None and n.is_integer() and 1 <= n <= ultimo_dia:
                    dias_col.setdefault(int(n), c)
            c_mat_j = _col(cab_j, 'MATR')
            c_comp = _col(cab_j, 'COMPLEMENTO')
            for linha in just[i_j + 1:]:
                mat = normalizar_matricula(_cel(linha, c_mat_j))
                if not mat:
                    continue
                dias = {}
                for dia, c in dias_col.items():
                    sigla = _txt(_cel(linha, c))
                    if sigla:
                        dias[str(dia)] = sigla.upper()
                complemento = _txt(_cel(linha, c_comp))
                if dias or complemento or mat in lancamentos:
                    l = _lanc(mat)
                    l['dias'] = dias
                    l['complemento'] = complemento

    horarios = _aba(abas, 'HORARIOS')
    if horarios:
        i_h, cab_h = _achar_cabecalho(horarios, ['MAT', 'HORARIO'])
        if i_h is not None:
            c_mat_h = _col(cab_h, 'MAT')
            c_hor = _col(cab_h, 'HORARIO')
            for linha in horarios[i_h + 1:]:
                mat = normalizar_matricula(_cel(linha, c_mat_h))
                hor = _txt(_cel(linha, c_hor))
                if mat and hor:
                    _lanc(mat)['horario'] = hor

    banco = _aba(abas, 'BANCO DE HORAS')
    if banco:
        i_b, cab_b = _achar_cabecalho(banco, ['MAT', 'SALDO ANTERIOR'])
        if i_b is not None:
            c_mat_b = _col(cab_b, 'MAT')
            c_ant = _col(cab_b, 'SALDO ANTERIOR')
            c_real = _col(cab_b, 'HORAS REALIZADAS')
            c_util = _col(cab_b, 'HORAS UTILIZADAS')
            c_atual = _col(cab_b, 'SALDO ATUAL')
            c_obs = _col(cab_b, 'OBSERVACOES', 'OBS')
            for linha in banco[i_b + 1:]:
                mat = normalizar_matricula(_cel(linha, c_mat_b))
                if not mat:
                    continue
                atual = _num(_cel(linha, c_atual))
                anterior = _num(_cel(linha, c_ant))
                if atual is None and anterior is None:
                    continue
                l = _lanc(mat)
                l['banco_saldo_anterior'] = anterior
                l['banco_realizadas'] = _num(_cel(linha, c_real))
                l['banco_utilizadas'] = _num(_cel(linha, c_util))
                l['banco_saldo_atual'] = atual
                l['banco_obs'] = _txt(_cel(linha, c_obs))

    horas_extras = []
    he = _aba(abas, 'HORAS EXTRAS')
    if he:
        i_e, cab_e = _achar_cabecalho(he, ['MAT', 'HE 50%'])
        if i_e is not None:
            c_mat_e = _col(cab_e, 'MAT')
            cols = {
                'he50': _col(cab_e, 'HE 50%'),
                'he100': _col(cab_e, 'HE 100%'),
                'he50_not': _col(cab_e, 'HE 50% NOT'),
                'he100_not': _col(cab_e, 'HE 100% NOT'),
            }
            # 'HE 50%' casa também com 'HE 50% NOT' por prefixo: garante a coluna exata.
            cols['he50'] = cab_e.get('HE 50%', cols['he50'])
            cols['he100'] = cab_e.get('HE 100%', cols['he100'])
            for linha in he[i_e + 1:]:
                mat = normalizar_matricula(_cel(linha, c_mat_e))
                if not mat:
                    continue
                valores = {k: (_num(_cel(linha, c)) or 0.0) for k, c in cols.items()}
                if not any(valores.values()):
                    continue
                horas_extras.append({'matricula': mat, 'nome': None, 'funcao': None, **valores})

    servidores = []
    bd = _aba(abas, 'BANCO DE DADOS')
    if bd:
        for linha in bd:
            mat = normalizar_matricula(_cel(linha, 0))
            nome = _txt(_cel(linha, 1))
            if mat and nome and not _norm(nome).startswith('NOME'):
                servidores.append({
                    'matricula': mat,
                    'nome': nome.upper(),
                    'funcao': _txt(_cel(linha, 2)),
                    'local': _txt(_cel(linha, 3)),
                })

    por_mat = {s['matricula']: s for s in servidores}
    for l in lancamentos.values():
        base = por_mat.get(l['matricula'])
        if base:
            l['nome'] = l['nome'] or base['nome']
            l['funcao'] = l['funcao'] or base['funcao']
    for h in horas_extras:
        origem = lancamentos.get(h['matricula']) or por_mat.get(h['matricula']) or {}
        h['nome'] = origem.get('nome')
        h['funcao'] = origem.get('funcao')

    for l in lancamentos.values():
        if l['banco_saldo_atual'] is None:
            continue
        esperado = (l['banco_saldo_anterior'] or 0) + (l['banco_realizadas'] or 0) - (l['banco_utilizadas'] or 0)
        if abs(esperado - l['banco_saldo_atual']) >= 0.01:
            avisos.append(
                f'Banco de horas de {l["nome"] or l["matricula"]} ({l["matricula"]}): '
                f'saldo atual {l["banco_saldo_atual"]:g} h, mas a conta dá {esperado:g} h.'
            )

    if not lancamentos:
        raise PlanilhaInvalida('A aba "Capa" não tem nenhuma matrícula preenchida.')

    for registro in [*lancamentos.values(), *horas_extras, *servidores]:
        for campo, limite in _LIMITES.items():
            if isinstance(registro.get(campo), str):
                registro[campo] = registro[campo][:limite]

    return {
        'arquivo': (nome_arquivo or '')[:255],
        'local': local[:120],
        'competencia': competencia.isoformat(),
        'lancamentos': sorted(lancamentos.values(), key=lambda x: (x['nome'] or '', x['matricula'])),
        'horas_extras': horas_extras,
        'servidores': servidores,
        'avisos': avisos,
    }
