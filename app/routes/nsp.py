# -*- coding: utf-8 -*-
"""Segurança do Paciente — notificações internas do NSP da unidade."""
import csv
import io
import os
import re
import uuid
import mimetypes
from collections import Counter
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from flask import (
    Blueprint, render_template, redirect, url_for, flash, request, abort,
    Response, send_from_directory, jsonify,
)
from flask_login import login_required, current_user
from sqlalchemy import func, or_
from sqlalchemy.orm import joinedload

from app import db
from app.models.nsp import (
    NspCatalogo, NspOcorrencia, NspAnexo, NspAndamento, NspEncaminhamento, NspAcao,
    GRUPOS_CATALOGO, GRUPOS_LABELS, BADGE_CORES,
)
from app.models.unidade import Unidade, UsuarioUnidade
from app.models.usuario import Usuario
from app.models.notificacao import Notificacao
from app.utils import agora_local

nsp_bp = Blueprint('nsp', __name__, url_prefix='/seguranca-paciente')

_UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'static', 'uploads', 'nsp')
_MAX_MB = 10
_HEADER_FILL = PatternFill('solid', fgColor='0D3B5E')
_HEADER_FONT = Font(bold=True, color='FFFFFF')
_HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _exigir(*acoes):
    if not any(current_user.pode(a) for a in acoes):
        abort(403)


def _unidades_visiveis():
    if current_user.perfil in ('administrador', 'gestor_secretaria') or current_user.pode('ver_todas_unidades'):
        return Unidade.query.filter_by(status='ativa').order_by(Unidade.nome).all()
    return current_user.unidades_ativas


def _ids_unidades_visiveis():
    return [u.id for u in _unidades_visiveis()]


def _pode_ver(ocorrencia):
    if current_user.perfil in ('administrador', 'gestor_secretaria'):
        return True
    return ocorrencia.unidade_id in _ids_unidades_visiveis()


def _ocorrencia_ou_404(id):
    o = NspOcorrencia.query.get_or_404(id)
    if not current_user.pode('ver_nsp') and not current_user.pode('adicionar_nsp') and not current_user.pode('editar_nsp'):
        abort(403)
    if not _pode_ver(o):
        abort(403)
    return o


def _catalogos_form():
    return {grupo: NspCatalogo.ativos(grupo) for grupo, _ in GRUPOS_CATALOGO}


def _gerar_protocolo():
    ano = datetime.utcnow().year
    prefixo = f'SP-{ano}-'
    ultimo = (
        db.session.query(func.max(NspOcorrencia.protocolo))
        .filter(NspOcorrencia.protocolo.like(f'{prefixo}%'))
        .scalar()
    )
    n = 1
    if ultimo:
        try:
            n = int(str(ultimo).split('-')[-1]) + 1
        except (TypeError, ValueError):
            n = 1
    return f'{prefixo}{n:05d}'


def _parse_date(valor):
    valor = (valor or '').strip()
    if not valor:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(valor, fmt).date()
        except ValueError:
            continue
    return None


def _parse_time(valor):
    valor = (valor or '').strip()
    if not valor:
        return None
    for fmt in ('%H:%M', '%H:%M:%S'):
        try:
            return datetime.strptime(valor, fmt).time()
        except ValueError:
            continue
    return None


def _int_or_none(valor):
    try:
        v = int(valor)
        return v if v else None
    except (TypeError, ValueError):
        return None


def _cat_id(form, campo):
    return _int_or_none(form.get(campo))


def _registrar_andamento(ocorrencia, tipo, texto, commit=False):
    andamento = NspAndamento(
        ocorrencia_id=ocorrencia.id,
        tipo=tipo,
        texto=texto,
        criado_por=current_user.id,
    )
    db.session.add(andamento)
    ocorrencia.atualizado_em = datetime.utcnow()
    if commit:
        db.session.commit()
    return andamento


def _notificar_unidade(unidade_id, tipo, titulo, texto, ocorrencia, excluir_id=None):
    if not unidade_id:
        return
    q = (
        db.session.query(Usuario.id)
        .join(UsuarioUnidade, UsuarioUnidade.usuario_id == Usuario.id)
        .filter(
            UsuarioUnidade.unidade_id == unidade_id,
            UsuarioUnidade.ativo == True,  # noqa: E712
            Usuario.ativo == True,  # noqa: E712
        )
    )
    ids = {row[0] for row in q.all()}
    extra = Usuario.query.filter(
        Usuario.perfil.in_(['administrador', 'gestor_secretaria']),
        Usuario.ativo == True,  # noqa: E712
    ).all()
    ids.update(u.id for u in extra)
    if excluir_id:
        ids.discard(excluir_id)
    for uid in ids:
        db.session.add(Notificacao(
            usuario_id=uid,
            tipo=tipo,
            titulo=titulo,
            texto=texto,
            nsp_ocorrencia_id=ocorrencia.id,
        ))


def _salvar_anexos(ocorrencia, files):
    os.makedirs(_UPLOAD_DIR, exist_ok=True)
    gravados = 0
    for f in files:
        if not f or not getattr(f, 'filename', None):
            continue
        f.stream.seek(0, os.SEEK_END)
        tamanho = f.stream.tell()
        f.stream.seek(0)
        if tamanho > _MAX_MB * 1024 * 1024:
            flash(f'Arquivo "{f.filename}" ultrapassa {_MAX_MB} MB e foi ignorado.', 'warning')
            continue
        original = f.filename
        mime = f.mimetype or mimetypes.guess_type(original)[0] or 'application/octet-stream'
        ext = os.path.splitext(original)[1].lower() or ''
        filename = f'{uuid.uuid4().hex}{ext}'
        f.save(os.path.join(_UPLOAD_DIR, filename))
        db.session.add(NspAnexo(
            ocorrencia_id=ocorrencia.id,
            filename=filename,
            original=original,
            mime_type=mime,
            tamanho_bytes=tamanho,
            criado_por=current_user.id,
        ))
        gravados += 1
    return gravados


def _query_base():
    q = NspOcorrencia.query.options(
        joinedload(NspOcorrencia.unidade),
        joinedload(NspOcorrencia.status),
        joinedload(NspOcorrencia.classificacao),
        joinedload(NspOcorrencia.tipo_incidente),
        joinedload(NspOcorrencia.setor_ocorrencia),
        joinedload(NspOcorrencia.departamento),
        joinedload(NspOcorrencia.autor),
    )
    ids = _ids_unidades_visiveis()
    if current_user.perfil not in ('administrador', 'gestor_secretaria') and not current_user.pode('ver_todas_unidades'):
        if not ids:
            return q.filter(False)
        q = q.filter(NspOcorrencia.unidade_id.in_(ids))
    return q


def _aplicar_filtros(q):
    unidade_ids = request.args.getlist('unidade', type=int)
    status_ids = request.args.getlist('status', type=int)
    class_ids = request.args.getlist('classificacao', type=int)
    tipo_ids = request.args.getlist('tipo_incidente', type=int)
    protocolo = (request.args.get('protocolo') or '').strip()
    busca = (request.args.get('q') or '').strip()
    never = request.args.get('never_event')
    dt_ini = _parse_date(request.args.get('de'))
    dt_fim = _parse_date(request.args.get('ate'))
    abertas = request.args.get('abertas')

    if unidade_ids:
        visiveis = set(_ids_unidades_visiveis())
        unidade_ids = [i for i in unidade_ids if i in visiveis] if visiveis or current_user.perfil in ('administrador', 'gestor_secretaria') else []
        if unidade_ids:
            q = q.filter(NspOcorrencia.unidade_id.in_(unidade_ids))
    if status_ids:
        q = q.filter(NspOcorrencia.status_id.in_(status_ids))
    if class_ids:
        q = q.filter(NspOcorrencia.classificacao_id.in_(class_ids))
    if tipo_ids:
        q = q.filter(NspOcorrencia.tipo_incidente_id.in_(tipo_ids))
    if protocolo:
        q = q.filter(NspOcorrencia.protocolo.ilike(f'%{protocolo}%'))
    if busca:
        like = f'%{busca}%'
        q = q.filter(or_(
            NspOcorrencia.descricao.ilike(like),
            NspOcorrencia.nome_afetado.ilike(like),
            NspOcorrencia.notificante_nome.ilike(like),
            NspOcorrencia.protocolo.ilike(like),
        ))
    if never == '1':
        q = q.filter(NspOcorrencia.never_event.is_(True))
    if dt_ini:
        q = q.filter(NspOcorrencia.data_ocorrencia >= dt_ini)
    if dt_fim:
        q = q.filter(NspOcorrencia.data_ocorrencia <= dt_fim)
    if abertas == '1':
        q = q.join(NspCatalogo, NspOcorrencia.status_id == NspCatalogo.id).filter(NspCatalogo.encerra.is_(False))
    return q


def _totais(q_base):
    itens = q_base.all()
    abertas = sum(1 for o in itens if not o.encerrada)
    graves = sum(1 for o in itens if o.classificacao and o.classificacao.slug in ('dano_grave', 'obito'))
    never = sum(1 for o in itens if o.never_event)
    encaminhadas = sum(1 for o in itens if o.status and o.status.slug == 'encaminhado')
    return {
        'total': len(itens),
        'abertas': abertas,
        'graves': graves,
        'never': never,
        'encaminhadas': encaminhadas,
    }


# ══════════════════════════════════════════════════════════
#  LISTAGEM / DASHBOARD
# ══════════════════════════════════════════════════════════

@nsp_bp.route('/')
@login_required
def index():
    _exigir('ver_nsp', 'adicionar_nsp', 'editar_nsp')
    q = _aplicar_filtros(_query_base())
    ocorrencias = q.order_by(NspOcorrencia.criado_em.desc()).limit(300).all()
    totais = _totais(_query_base())
    return render_template(
        'nsp/index.html',
        ocorrencias=ocorrencias,
        totais=totais,
        unidades=_unidades_visiveis(),
        status_opts=NspCatalogo.ativos('status'),
        classificacoes=NspCatalogo.ativos('classificacao'),
        tipos_incidente=NspCatalogo.ativos('tipo_incidente'),
        filtros=request.args,
    )


def _ctx_form(unidades, unidade_padrao_id, form=None):
    from app.sis_consulta import configurado as sis_configurado
    paciente = NspCatalogo.por_slug('tipo_pessoa', 'paciente')
    return dict(
        unidades=unidades,
        catalogos=_catalogos_form(),
        unidade_padrao_id=unidade_padrao_id,
        form=form,
        sis_ok=sis_configurado(),
        tipo_pessoa_paciente_id=paciente.id if paciente else None,
    )


# ══════════════════════════════════════════════════════════
#  NOVA NOTIFICAÇÃO
# ══════════════════════════════════════════════════════════

@nsp_bp.route('/nova', methods=['GET', 'POST'])
@login_required
def nova():
    _exigir('adicionar_nsp')
    unidades = _unidades_visiveis()
    if not unidades:
        flash('Você precisa estar vinculado a uma unidade para notificar.', 'warning')
        return redirect(url_for('nsp.index'))

    unidade_padrao_id = request.args.get('unidade_id', type=int)
    if not unidade_padrao_id and current_user.unidade_logada:
        unidade_padrao_id = current_user.unidade_logada.id

    if request.method == 'POST':
        unidade_id = _int_or_none(request.form.get('unidade_id'))
        if unidade_id not in {u.id for u in unidades}:
            flash('Unidade inválida.', 'danger')
            return redirect(url_for('nsp.nova'))

        notificante_nome = (request.form.get('notificante_nome') or '').strip() or current_user.nome
        descricao = (request.form.get('descricao') or '').strip()
        acao_imediata = (request.form.get('acao_imediata') or '').strip()
        data_ocorrencia = _parse_date(request.form.get('data_ocorrencia'))
        classificacao_id = _cat_id(request.form, 'classificacao_id')

        if not descricao or not acao_imediata or not data_ocorrencia or not classificacao_id:
            flash('Preencha unidade, data, classificação, descrição e ação imediata.', 'danger')
            return render_template('nsp/form.html', **_ctx_form(unidades, unidade_id, request.form))

        tipo_incidente = NspCatalogo.query.get(_cat_id(request.form, 'tipo_incidente_id')) if _cat_id(request.form, 'tipo_incidente_id') else None
        never = request.form.get('never_event') == '1' or bool(tipo_incidente and tipo_incidente.never_event)
        status_aberto = NspCatalogo.por_slug('status', 'aberto')

        o = NspOcorrencia(
            protocolo=_gerar_protocolo(),
            unidade_id=unidade_id,
            notificante_nome=notificante_nome,
            tipo_setor_notificante_id=_cat_id(request.form, 'tipo_setor_notificante_id'),
            setor_notificante_id=_cat_id(request.form, 'setor_notificante_id'),
            setor_notificante_outro=(request.form.get('setor_notificante_outro') or '').strip() or None,
            nome_afetado=(request.form.get('nome_afetado') or '').strip() or None,
            data_nascimento=_parse_date(request.form.get('data_nascimento')),
            prontuario=(request.form.get('prontuario') or '').strip() or None,
            cpf=re.sub(r'\D+', '', request.form.get('cpf') or '') or None,
            cns=re.sub(r'\D+', '', request.form.get('cns') or '') or None,
            tipo_pessoa_id=_cat_id(request.form, 'tipo_pessoa_id'),
            tipo_pessoa_outro=(request.form.get('tipo_pessoa_outro') or '').strip() or None,
            data_ocorrencia=data_ocorrencia,
            hora_ocorrencia=_parse_time(request.form.get('hora_ocorrencia')),
            tipo_setor_ocorrencia_id=_cat_id(request.form, 'tipo_setor_ocorrencia_id'),
            setor_ocorrencia_id=_cat_id(request.form, 'setor_ocorrencia_id'),
            setor_ocorrencia_outro=(request.form.get('setor_ocorrencia_outro') or '').strip() or None,
            departamento_id=_cat_id(request.form, 'departamento_id'),
            tipo_incidente_id=_cat_id(request.form, 'tipo_incidente_id'),
            classificacao_id=classificacao_id,
            never_event=never,
            descricao=descricao,
            acao_imediata=acao_imediata,
            tipo_setor_notificado_id=_cat_id(request.form, 'tipo_setor_notificado_id'),
            setor_notificado_id=_cat_id(request.form, 'setor_notificado_id'),
            setor_notificado_outro=(request.form.get('setor_notificado_outro') or '').strip() or None,
            status_id=status_aberto.id if status_aberto else None,
            criado_por=current_user.id,
        )
        db.session.add(o)
        db.session.flush()
        _salvar_anexos(o, request.files.getlist('anexos'))
        _registrar_andamento(o, 'registro', f'Notificação {o.protocolo} aberta por {current_user.nome}.')
        _notificar_unidade(
            o.unidade_id, 'nsp_novo',
            f'Segurança do paciente {o.protocolo}',
            f'Nova ocorrência na unidade {o.unidade.nome if o.unidade else ""}.',
            o, excluir_id=current_user.id,
        )
        db.session.commit()
        flash(f'Notificação {o.protocolo} registrada.', 'success')
        return redirect(url_for('nsp.detalhe', id=o.id))

    return render_template('nsp/form.html', **_ctx_form(unidades, unidade_padrao_id))


@nsp_bp.route('/buscar-sis', methods=['POST'])
@login_required
def buscar_sis():
    _exigir('adicionar_nsp', 'editar_nsp', 'ver_nsp')
    from app.sis_consulta import buscar_paciente
    payload = request.get_json(silent=True) or {}
    termo = (request.form.get('q') or payload.get('q') or '').strip()
    tipo = (request.form.get('tipo') or payload.get('tipo') or '').strip().lower()
    if not termo:
        return jsonify(ok=False, erro='Informe o CPF, o CNS ou o prontuário.'), 400
    try:
        dados = buscar_paciente(termo, tipo=tipo or None)
        return jsonify(dados)
    except ValueError as exc:
        return jsonify(ok=False, erro=str(exc)), 400
    except Exception as exc:
        return jsonify(ok=False, erro=str(exc) or 'Não foi possível consultar o SIS.'), 502


# ══════════════════════════════════════════════════════════
#  DETALHE / ACOMPANHAMENTO
# ══════════════════════════════════════════════════════════

@nsp_bp.route('/<int:id>')
@login_required
def detalhe(id):
    _exigir('ver_nsp', 'adicionar_nsp', 'editar_nsp')
    o = _ocorrencia_ou_404(id)
    return render_template(
        'nsp/detalhe.html',
        o=o,
        catalogos=_catalogos_form(),
        unidades=_unidades_visiveis(),
        pode_editar=current_user.pode('editar_nsp') and not o.encerrada,
        pode_comentar=current_user.pode('ver_nsp') or current_user.pode('editar_nsp') or current_user.pode('adicionar_nsp'),
    )


@nsp_bp.route('/<int:id>/imprimir')
@login_required
def imprimir(id):
    _exigir('ver_nsp', 'adicionar_nsp', 'editar_nsp')
    o = _ocorrencia_ou_404(id)
    fotos_img = [a for a in o.anexos if a.is_image]
    return render_template('nsp/imprimir.html', o=o, now=agora_local(), fotos_img=fotos_img)


@nsp_bp.route('/<int:id>/comentario', methods=['POST'])
@login_required
def comentario(id):
    o = _ocorrencia_ou_404(id)
    if o.encerrada:
        flash('Ocorrência encerrada.', 'warning')
        return redirect(url_for('nsp.detalhe', id=o.id))
    texto = (request.form.get('texto') or '').strip()
    if not texto:
        flash('Escreva o andamento.', 'danger')
        return redirect(url_for('nsp.detalhe', id=o.id))
    _registrar_andamento(o, 'comentario', texto)
    _notificar_unidade(
        o.unidade_id, 'nsp_andamento',
        f'Andamento em {o.protocolo}',
        texto[:180],
        o, excluir_id=current_user.id,
    )
    db.session.commit()
    flash('Andamento registrado.', 'success')
    return redirect(url_for('nsp.detalhe', id=o.id))


@nsp_bp.route('/<int:id>/status', methods=['POST'])
@login_required
def mudar_status(id):
    _exigir('editar_nsp')
    o = _ocorrencia_ou_404(id)
    status = NspCatalogo.query.get(_cat_id(request.form, 'status_id'))
    if not status or status.grupo != 'status':
        flash('Status inválido.', 'danger')
        return redirect(url_for('nsp.detalhe', id=o.id))
    if status.encerra and o.exige_investigacao and not o.investigacao_preenchida:
        flash('Óbito, dano grave e evento que nunca deveria ocorrer exigem as 6 etapas da investigação antes de encerrar.', 'danger')
        return redirect(url_for('nsp.detalhe', id=o.id))
    anterior = o.status_nome
    o.status_id = status.id
    o.encerrado_em = datetime.utcnow() if status.encerra else None
    _registrar_andamento(o, 'status', f'Status alterado de {anterior} para {status.nome}.')
    db.session.commit()
    flash(f'Status atualizado para {status.nome}.', 'success')
    return redirect(url_for('nsp.detalhe', id=o.id))


@nsp_bp.route('/<int:id>/analise', methods=['POST'])
@login_required
def salvar_analise(id):
    _exigir('editar_nsp')
    o = _ocorrencia_ou_404(id)
    if o.encerrada:
        flash('Ocorrência encerrada.', 'warning')
        return redirect(url_for('nsp.detalhe', id=o.id))
    o.fatores_contribuintes = (request.form.get('fatores_contribuintes') or '').strip() or None
    o.consequencias_organizacionais = (request.form.get('consequencias_organizacionais') or '').strip() or None
    o.deteccao = (request.form.get('deteccao') or '').strip() or None
    o.fatores_atenuantes = (request.form.get('fatores_atenuantes') or '').strip() or None
    o.acoes_melhoria = (request.form.get('acoes_melhoria') or '').strip() or None
    o.acoes_reducao_risco = (request.form.get('acoes_reducao_risco') or '').strip() or None
    o.analise_resumo = (request.form.get('analise_resumo') or '').strip() or None
    o.analise_por = current_user.id
    o.analise_em = datetime.utcnow()
    o.never_event = request.form.get('never_event') == '1'
    o.notivisa_notificado = request.form.get('notivisa_notificado') == '1'
    o.notivisa_numero = (request.form.get('notivisa_numero') or '').strip() or None
    o.notivisa_em = _parse_date(request.form.get('notivisa_em'))
    status_analise = NspCatalogo.por_slug('status', 'em_analise')
    if status_analise and o.status and o.status.slug == 'aberto':
        o.status_id = status_analise.id
    _registrar_andamento(o, 'analise', 'Investigação / análise atualizada (etapas da Anvisa).')
    if o.notivisa_notificado:
        _registrar_andamento(
            o, 'notivisa',
            f'Registro no Notivisa: {o.notivisa_numero or "sem número"}.'
        )
    db.session.commit()
    flash('Análise salva.', 'success')
    return redirect(url_for('nsp.detalhe', id=o.id))


@nsp_bp.route('/<int:id>/encaminhar', methods=['POST'])
@login_required
def encaminhar(id):
    _exigir('editar_nsp')
    o = _ocorrencia_ou_404(id)
    if o.encerrada:
        flash('Ocorrência encerrada.', 'warning')
        return redirect(url_for('nsp.detalhe', id=o.id))
    destino_id = _cat_id(request.form, 'destino_id')
    if not destino_id:
        flash('Escolha o setor de destino.', 'danger')
        return redirect(url_for('nsp.detalhe', id=o.id))
    unidade_destino_id = _int_or_none(request.form.get('unidade_destino_id')) or o.unidade_id
    texto = (request.form.get('texto') or '').strip()
    enc = NspEncaminhamento(
        ocorrencia_id=o.id,
        destino_id=destino_id,
        destino_outro=(request.form.get('destino_outro') or '').strip() or None,
        unidade_destino_id=unidade_destino_id,
        texto=texto or None,
        criado_por=current_user.id,
    )
    db.session.add(enc)
    db.session.flush()
    status_enc = NspCatalogo.por_slug('status', 'encaminhado')
    if status_enc:
        o.status_id = status_enc.id
    dest_nome = enc.destino_nome
    un_nome = enc.unidade_destino.nome if enc.unidade_destino else o.unidade.nome
    _registrar_andamento(
        o, 'encaminhamento',
        f'Encaminhado para {dest_nome} ({un_nome}). {texto}'.strip(),
    )
    _notificar_unidade(
        unidade_destino_id, 'nsp_encaminhado',
        f'{o.protocolo} encaminhada',
        f'Encaminhada para {dest_nome}. {texto or ""}'.strip(),
        o, excluir_id=current_user.id,
    )
    db.session.commit()
    flash(f'Ocorrência encaminhada para {dest_nome}.', 'success')
    return redirect(url_for('nsp.detalhe', id=o.id))


@nsp_bp.route('/<int:id>/acao', methods=['POST'])
@login_required
def nova_acao(id):
    _exigir('editar_nsp')
    o = _ocorrencia_ou_404(id)
    if o.encerrada:
        flash('Ocorrência encerrada.', 'warning')
        return redirect(url_for('nsp.detalhe', id=o.id))
    descricao = (request.form.get('descricao') or '').strip()
    if not descricao:
        flash('Descreva a ação.', 'danger')
        return redirect(url_for('nsp.detalhe', id=o.id))
    acao = NspAcao(
        ocorrencia_id=o.id,
        descricao=descricao,
        responsavel_nome=(request.form.get('responsavel_nome') or '').strip() or None,
        prazo=_parse_date(request.form.get('prazo')),
        criado_por=current_user.id,
    )
    db.session.add(acao)
    status_plano = NspCatalogo.por_slug('status', 'plano_acao')
    if status_plano and o.status and o.status.slug in ('aberto', 'em_analise', 'encaminhado'):
        o.status_id = status_plano.id
    _registrar_andamento(o, 'acao', f'Ação incluída: {descricao[:180]}')
    db.session.commit()
    flash('Ação incluída no plano.', 'success')
    return redirect(url_for('nsp.detalhe', id=o.id))


@nsp_bp.route('/<int:id>/acao/<int:acao_id>/concluir', methods=['POST'])
@login_required
def concluir_acao(id, acao_id):
    _exigir('editar_nsp')
    o = _ocorrencia_ou_404(id)
    acao = NspAcao.query.filter_by(id=acao_id, ocorrencia_id=o.id).first_or_404()
    acao.concluida = not acao.concluida
    acao.concluida_em = datetime.utcnow() if acao.concluida else None
    _registrar_andamento(
        o, 'acao',
        f'Ação {"concluída" if acao.concluida else "reaberta"}: {acao.descricao[:160]}',
    )
    db.session.commit()
    return redirect(url_for('nsp.detalhe', id=o.id))


@nsp_bp.route('/<int:id>/anexos', methods=['POST'])
@login_required
def anexar(id):
    o = _ocorrencia_ou_404(id)
    if not (current_user.pode('editar_nsp') or current_user.pode('adicionar_nsp')):
        abort(403)
    n = _salvar_anexos(o, request.files.getlist('anexos'))
    if n:
        _registrar_andamento(o, 'anexo', f'{n} arquivo(s) anexado(s).')
        db.session.commit()
        flash(f'{n} arquivo(s) anexado(s).', 'success')
    else:
        flash('Nenhum arquivo válido enviado (limite 10 MB).', 'warning')
    return redirect(url_for('nsp.detalhe', id=o.id))


@nsp_bp.route('/anexo/<int:id>/baixar')
@login_required
def baixar_anexo(id):
    anexo = NspAnexo.query.get_or_404(id)
    _ocorrencia_ou_404(anexo.ocorrencia_id)
    return send_from_directory(_UPLOAD_DIR, anexo.filename, as_attachment=True,
                               download_name=anexo.original or anexo.filename)


# ══════════════════════════════════════════════════════════
#  CONSULTA DE PROTOCOLO
# ══════════════════════════════════════════════════════════

@nsp_bp.route('/consultar')
@login_required
def consultar():
    _exigir('ver_nsp', 'adicionar_nsp', 'editar_nsp')
    protocolo = (request.args.get('protocolo') or '').strip().upper()
    o = None
    if protocolo:
        o = NspOcorrencia.query.filter(func.upper(NspOcorrencia.protocolo) == protocolo).first()
        if o and not _pode_ver(o):
            o = None
            flash('Protocolo não encontrado nesta unidade.', 'warning')
        elif o:
            return redirect(url_for('nsp.detalhe', id=o.id))
        else:
            flash('Protocolo não encontrado.', 'warning')
    return render_template('nsp/consultar.html', protocolo=protocolo)


# ══════════════════════════════════════════════════════════
#  RELATÓRIOS
# ══════════════════════════════════════════════════════════

def _linhas_exportacao(ocorrencias):
    cab = [
        'Protocolo', 'Unidade', 'Data ocorrência', 'Hora', 'Status', 'Classificação',
        'Tipo de incidente', 'Jamais deveria ocorrer', 'Setor', 'Departamento', 'Tipo pessoa',
        'Pessoa afetada', 'Prontuário', 'CPF', 'CNS', 'Notificante', 'Descrição', 'Ação imediata',
        'Notivisa', 'Nº Notivisa', 'Aberto em',
    ]
    linhas = [cab]
    for o in ocorrencias:
        linhas.append([
            o.protocolo,
            o.unidade.nome if o.unidade else '',
            o.data_ocorrencia.strftime('%d/%m/%Y') if o.data_ocorrencia else '',
            o.hora_ocorrencia.strftime('%H:%M') if o.hora_ocorrencia else '',
            o.status_nome,
            o.classificacao_nome,
            o.tipo_incidente_nome,
            'Sim' if o.never_event else 'Não',
            o.setor_ocorrencia_nome,
            o.departamento.nome if o.departamento else '',
            o.tipo_pessoa.nome if o.tipo_pessoa else '',
            o.nome_afetado or '',
            o.prontuario or '',
            o.cpf or '',
            o.cns or '',
            o.notificante_nome,
            (o.descricao or '').replace('\n', ' '),
            (o.acao_imediata or '').replace('\n', ' '),
            'Sim' if o.notivisa_notificado else 'Não',
            o.notivisa_numero or '',
            o.criado_em.strftime('%d/%m/%Y %H:%M') if o.criado_em else '',
        ])
    return linhas


@nsp_bp.route('/relatorios')
@login_required
def relatorios():
    _exigir('ver_nsp', 'editar_nsp')
    q = _aplicar_filtros(_query_base())
    ocorrencias = q.order_by(NspOcorrencia.data_ocorrencia.desc()).all()

    por_status = Counter((o.status_nome, o.status_cor) for o in ocorrencias)
    por_class = Counter((o.classificacao_nome, o.classificacao_cor) for o in ocorrencias)
    por_tipo = Counter(o.tipo_incidente_nome for o in ocorrencias)
    por_unidade = Counter((o.unidade.nome if o.unidade else '—') for o in ocorrencias)
    por_mes = Counter()
    for o in ocorrencias:
        if o.data_ocorrencia:
            por_mes[o.data_ocorrencia.strftime('%Y-%m')] += 1

    return render_template(
        'nsp/relatorios.html',
        ocorrencias=ocorrencias,
        totais=_totais(_query_base()),
        por_status=sorted(por_status.items(), key=lambda x: -x[1]),
        por_class=sorted(por_class.items(), key=lambda x: -x[1]),
        por_tipo=por_tipo.most_common(12),
        por_unidade=por_unidade.most_common(15),
        por_mes=sorted(por_mes.items()),
        unidades=_unidades_visiveis(),
        status_opts=NspCatalogo.ativos('status'),
        classificacoes=NspCatalogo.ativos('classificacao'),
        tipos_incidente=NspCatalogo.ativos('tipo_incidente'),
        filtros=request.args,
    )


@nsp_bp.route('/relatorios/exportar/<formato>')
@login_required
def exportar(formato):
    _exigir('ver_nsp', 'editar_nsp')
    q = _aplicar_filtros(_query_base())
    ocorrencias = q.order_by(NspOcorrencia.data_ocorrencia.desc()).all()
    linhas = _linhas_exportacao(ocorrencias)
    ts = datetime.now().strftime('%Y%m%d_%H%M')
    nome = f'seguranca_paciente_{ts}'

    if formato == 'csv':
        buf = io.StringIO()
        writer = csv.writer(buf, delimiter=';')
        writer.writerows(linhas)
        dados = '\ufeff' + buf.getvalue()
        return Response(
            dados,
            mimetype='text/csv; charset=utf-8',
            headers={'Content-Disposition': f'attachment; filename="{nome}.csv"'},
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Ocorrências'
    ws.append(linhas[0])
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
    for linha in linhas[1:]:
        ws.append(linha)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(42, max(12, len(str(col[0].value or '')) + 4))

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    ext = 'xls' if formato == 'xls' else 'xlsx'
    mime = 'application/vnd.ms-excel' if formato == 'xls' else \
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return Response(
        out.getvalue(),
        mimetype=mime,
        headers={'Content-Disposition': f'attachment; filename="{nome}.{ext}"'},
    )
