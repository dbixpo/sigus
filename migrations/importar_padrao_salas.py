# -*- coding: utf-8 -*-
"""Importa o padrão de ambientes, itens e kit por sala (planilha do Planejamento).

Uso:
    python migrations/importar_padrao_salas.py                 # simulação (não grava)
    python migrations/importar_padrao_salas.py --aplicar       # grava
    python migrations/importar_padrao_salas.py --planilha X.xlsx --pendencias saida.xlsx

Regras:
- Ambiente com equivalente no SIGUS (DE_PARA_SALAS): o tipo existente ganha o nome,
  o código AMB e o grupo da planilha; as salas continuam ligadas a ele.
- Ambiente sem equivalente: vira um tipo de sala novo.
- Tipos do SIGUS fora do catálogo ficam como estão (sem código).
- Item com equivalente (DE_PARA_ITENS): o tipo de equipamento mantém o nome e ganha
  código, classificação, valor de referência e BASE/FUNÇÃO. Os demais viram tipos novos.
- Kit: soma as linhas repetidas de ambiente + item; ignora linha sem código ou sem
  quantidade (vão para a planilha de pendências).
- Unidades: casadas pelo CNES; grava o código do imóvel.
Pode rodar de novo quando a planilha mudar: atualiza pelo código, sem duplicar.
"""
import argparse
import os
import sys
import warnings
from collections import OrderedDict, defaultdict
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from app import create_app, db
from app.models.equipamento import Equipamento, TipoEquipamento
from app.models.sala import Sala
from app.models.tipo_sala import KitPadraoSala, TipoSala
from app.models.unidade import Unidade

PLANILHA_PADRAO = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'dados', 'padrao_salas_ubs.xlsx')

DE_PARA_SALAS = {
    'AMB-01': 'Recepção',
    'AMB-04': 'Banheiro Feminino PCD',
    'AMB-05': 'Banheiro Masculino PCD',
    'AMB-13': 'Consultório de Ginecologia',
    'AMB-15': 'Acolhimento',
    'AMB-17': 'Consultório de Telessaúde',
    'AMB-21': 'Pós-consulta',
    'AMB-22': 'Consultório Odontológico',
    'AMB-24': 'Sala de Curativo',
    'AMB-26': 'Sala de Medicação',
    'AMB-27': 'Coleta',
    'AMB-29': 'Sala de Emergência',
    'AMB-30': 'Sala de Vacina',
    'AMB-33': 'Central de Material e Esterilização',
    'AMB-38': 'Farmácia',
    'AMB-40': 'Coordenação',
    'AMB-41': 'Sala de Reunião',
    'AMB-42': 'Sala de Agentes Comunitários',
    'AMB-45': 'Copa',
    'AMB-46': 'Vestiário Feminino',
    'AMB-47': 'Vestiário Masculino',
    'AMB-49': 'Almoxarifado',
}

DE_PARA_ITENS = {
    'ITEM-051': 'Computador (Gabinete)',
    'ITEM-052': 'Notebook',
    'ITEM-055': 'DEA (RCP)',
    'ITEM-105': 'Projetor',
    'ITEM-115': 'TV',
}

# Tipos do SIGUS sem código próprio que atendem a um item do catálogo.
CONTA_COMO = {
    'Computador All-In-One': 'ITEM-051',
}

ICONE_GRUPO = [
    ('RECEPÇÃO', 'fas fa-concierge-bell'),
    ('SAÚDE BUCAL', 'fas fa-tooth'),
    ('ASSISTENCIAL', 'fas fa-user-doctor'),
    ('APOIO TÉCNICO', 'fas fa-flask'),
    ('APOIO ADMINISTRATIVO', 'fas fa-briefcase'),
    ('APOIO AOS TRABALHADORES', 'fas fa-mug-hot'),
    ('APOIO LOGÍSTICO', 'fas fa-boxes-stacked'),
    ('RESÍDUOS', 'fas fa-recycle'),
    ('INFRAESTRUTURA TECNOLÓGICA', 'fas fa-network-wired'),
    ('INFRAESTRUTURA PREDIAL', 'fas fa-screwdriver-wrench'),
]

ICONE_CLASSIFICACAO = [
    ('Mobiliário', 'fas fa-chair'),
    ('informática', 'fas fa-desktop'),
    ('Equipamentos permanentes', 'fas fa-stethoscope'),
]


def limpar(valor):
    if valor is None:
        return ''
    return ' '.join(str(valor).split())


def icone_por(texto, tabela, padrao):
    for chave, icone in tabela:
        if chave.lower() in (texto or '').lower():
            return icone
    return padrao


def linhas_apos_cabecalho(ws, primeira_coluna):
    """(nº da linha, valores) das linhas abaixo do cabeçalho cuja 1ª célula é `primeira_coluna`."""
    achou = False
    for idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if not achou:
            achou = limpar(row[0]).upper() == primeira_coluna.upper()
            continue
        if any(v not in (None, '') for v in row):
            yield idx, row


def ler_planilha(caminho):
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        wb = openpyxl.load_workbook(caminho, data_only=True)
    pend = []

    ambientes = OrderedDict()
    for n, r in linhas_apos_cabecalho(wb['Ambientes Padrão'], 'Nº'):
        codigo = limpar(r[4])
        if not codigo.startswith('AMB-'):
            continue
        ambientes[codigo] = {
            'ordem': int(r[0]) if isinstance(r[0], (int, float)) else None,
            'grupo': limpar(r[1]),
            'nome': limpar(r[2]),
            'descricao': limpar(r[3]) or None,
        }

    itens = OrderedDict()
    for n, r in linhas_apos_cabecalho(wb['Itens'], 'CÓD.'):
        codigo = limpar(r[0])
        if not codigo.startswith('ITEM-'):
            continue
        classificacao = limpar(r[2])
        valor = r[3] if isinstance(r[3], (int, float)) else None
        natureza = {'BASE': 'base', 'FUNÇÃO': 'funcao'}.get(limpar(r[4]).upper())
        if classificacao in ('', '(a classificar)'):
            pend.append(('Itens', n, 'Item sem classificação', f'{codigo} {limpar(r[1])}'))
            classificacao = None
        if not valor:
            pend.append(('Itens', n, 'Item sem valor de referência', f'{codigo} {limpar(r[1])}'))
            valor = None
        if not natureza:
            pend.append(('Itens', n, 'Item sem BASE/FUNÇÃO', f'{codigo} {limpar(r[1])}'))
        itens[codigo] = {
            'nome': limpar(r[1]),
            'classificacao': classificacao,
            'valor': Decimal(str(valor)) if valor else None,
            'natureza': natureza,
        }

    kit = defaultdict(int)
    linhas_kit = defaultdict(list)
    for n, r in linhas_apos_cabecalho(wb['Kit — Lista'], 'CÓD. AMB.'):
        amb, item, nome_item, qtd = limpar(r[0]), limpar(r[3]), limpar(r[4]), r[5]
        if not item:
            pend.append(('Kit — Lista', n, 'Linha sem código de item (não importada)', f'{amb} · {nome_item}'))
            continue
        if item not in itens:
            pend.append(('Kit — Lista', n, 'Código de item que não existe na aba Itens', f'{amb} · {item} {nome_item}'))
            continue
        if amb not in ambientes:
            pend.append(('Kit — Lista', n, 'Código de ambiente que não existe na aba Ambientes', f'{amb} · {nome_item}'))
            continue
        if not isinstance(qtd, (int, float)):
            pend.append(('Kit — Lista', n, 'Linha sem quantidade (não importada)', f'{amb} · {item} {nome_item}'))
            continue
        if qtd <= 0:
            pend.append(('Kit — Lista', n, 'Quantidade zero (não importada)', f'{amb} · {item} {nome_item} · {limpar(r[8])}'))
            continue
        kit[(amb, item)] += int(qtd)
        linhas_kit[(amb, item)].append((n, int(qtd)))
    for (amb, item), ocorrencias in linhas_kit.items():
        if len(ocorrencias) > 1:
            detalhe = ' + '.join(f'{q} (linha {n})' for n, q in ocorrencias)
            pend.append(('Kit — Lista', ocorrencias[0][0], 'Item repetido no mesmo ambiente (quantidades somadas)',
                         f'{amb} · {item} {itens[item]["nome"]}: {detalhe} = {kit[(amb, item)]}'))
    for amb, a in ambientes.items():
        if not any(k[0] == amb for k in kit):
            pend.append(('Kit — Lista', '', 'Ambiente sem kit definido', f'{amb} {a["nome"]}'))

    unidades = []
    for n, r in linhas_apos_cabecalho(wb['Unidades'], 'Nº'):
        if not limpar(r[1]).startswith('IM-'):
            continue
        unidades.append({'linha': n, 'codigo_imovel': limpar(r[1]), 'nome': limpar(r[2]),
                         'cnes': limpar(r[3]), 'situacao': limpar(r[6])})

    return ambientes, itens, kit, unidades, pend


def importar(ambientes, itens, kit, unidades, pend, log):
    tipos_sala = {}
    for codigo, a in ambientes.items():
        tipo = TipoSala.query.filter_by(codigo=codigo).first()
        origem = 'código'
        if not tipo and codigo in DE_PARA_SALAS:
            tipo = TipoSala.query.filter_by(nome=DE_PARA_SALAS[codigo]).first()
            origem = f'de-para "{DE_PARA_SALAS[codigo]}"'
        if not tipo:
            tipo = TipoSala.query.filter_by(nome=a['nome']).first()
            origem = 'mesmo nome'
        if tipo:
            conflito = TipoSala.query.filter(TipoSala.nome == a['nome'], TipoSala.id != tipo.id).first()
            if conflito:
                pend.append(('SIGUS', '', 'Nome do ambiente já usado por outro tipo de sala',
                             f'{codigo} "{a["nome"]}" (tipo id {conflito.id}); mantido "{tipo.nome}"'))
            elif tipo.nome != a['nome']:
                log.append(('Tipo de sala renomeado', f'{tipo.nome} → {codigo} {a["nome"]} ({origem})'))
                tipo.nome = a['nome']
            else:
                log.append(('Tipo de sala casado', f'{codigo} {a["nome"]} ({origem})'))
        else:
            tipo = TipoSala(nome=a['nome'], icone=icone_por(a['grupo'], ICONE_GRUPO, 'fas fa-door-open'))
            db.session.add(tipo)
            log.append(('Tipo de sala criado', f'{codigo} {a["nome"]}'))
        tipo.codigo = codigo
        tipo.grupo = a['grupo'] or None
        tipo.ordem = a['ordem']
        tipo.ativo = True
        if not tipo.descricao and a['descricao']:
            tipo.descricao = a['descricao']
        if (tipo.icone or 'fas fa-door-open') in ('fas fa-door-open', 'bi bi-door-open', 'bi-door-open'):
            tipo.icone = icone_por(a['grupo'], ICONE_GRUPO, 'fas fa-door-open')
        tipos_sala[codigo] = tipo
    db.session.flush()

    tipos_equip = {}
    for codigo, it in itens.items():
        tipo = TipoEquipamento.query.filter_by(codigo=codigo).first()
        if not tipo and codigo in DE_PARA_ITENS:
            tipo = TipoEquipamento.query.filter_by(nome=DE_PARA_ITENS[codigo]).first()
            if tipo:
                log.append(('Tipo de equipamento casado', f'{tipo.nome} = {codigo} {it["nome"]}'))
        if not tipo:
            tipo = TipoEquipamento.query.filter_by(nome=it['nome']).first()
        if not tipo:
            tipo = TipoEquipamento(
                nome=it['nome'], tem_patrimonio=True, ativo=True,
                icone=icone_por(it['classificacao'], ICONE_CLASSIFICACAO, 'fas fa-box'),
            )
            db.session.add(tipo)
            log.append(('Tipo de equipamento criado', f'{codigo} {it["nome"]}'))
        tipo.codigo = codigo
        tipo.classificacao = it['classificacao']
        tipo.valor_referencia = it['valor']
        tipo.natureza_kit = it['natureza']
        tipos_equip[codigo] = tipo
    db.session.flush()

    for nome, codigo in CONTA_COMO.items():
        tipo = TipoEquipamento.query.filter_by(nome=nome).first()
        if tipo and codigo in tipos_equip and not tipo.conta_como_id:
            tipo.conta_como_id = tipos_equip[codigo].id
            log.append(('Tipo de equipamento equivalente', f'{nome} conta como {codigo} {tipos_equip[codigo].nome}'))

    novos = atualizados = 0
    for (amb, item), qtd in kit.items():
        ts, te = tipos_sala[amb], tipos_equip[item]
        linha = KitPadraoSala.query.filter_by(tipo_sala_id=ts.id, tipo_equipamento_id=te.id).first()
        if linha:
            if linha.quantidade != qtd:
                linha.quantidade = qtd
                atualizados += 1
        else:
            db.session.add(KitPadraoSala(tipo_sala_id=ts.id, tipo_equipamento_id=te.id, quantidade=qtd))
            novos += 1
    log.append(('Kit', f'{novos} linha(s) nova(s), {atualizados} quantidade(s) atualizada(s), '
                       f'{len(kit)} item(ns) de kit em {len({k[0] for k in kit})} ambientes'))

    for u in unidades:
        candidatas = Unidade.query.filter_by(numero_cnes=u['cnes']).all() if u['cnes'].isdigit() else []
        if len(candidatas) > 1:
            candidatas = [c for c in candidatas if (c.tipo or '').startswith('APS')] or candidatas
        if len(candidatas) != 1:
            motivo = 'não encontrada no SIGUS' if not candidatas else 'mais de uma unidade com o CNES'
            pend.append(('Unidades', u['linha'], f'Unidade {motivo}',
                         f'{u["codigo_imovel"]} {u["nome"]} (CNES {u["cnes"]}, {u["situacao"]})'))
            continue
        unidade = candidatas[0]
        if unidade.codigo_imovel != u['codigo_imovel']:
            unidade.codigo_imovel = u['codigo_imovel']
        log.append(('Unidade casada', f'{u["codigo_imovel"]} {u["nome"]} → {unidade.nome}'))
    db.session.flush()


def pendencias_sigus():
    """Tipos do SIGUS fora do padrão que ainda aparecem em uso."""
    salas = (db.session.query(TipoSala.nome, Unidade.tipo, db.func.count(Sala.id))
             .join(Sala, Sala.tipo_sala_id == TipoSala.id).join(Unidade, Sala.unidade_id == Unidade.id)
             .filter(TipoSala.codigo.is_(None), Sala.ativo.is_(True))
             .group_by(TipoSala.nome, Unidade.tipo).all())
    por_tipo = defaultdict(lambda: [0, 0])
    for nome, tipo_unidade, n in salas:
        por_tipo[nome][0 if (tipo_unidade or '').startswith('APS') else 1] += n
    salas_sem_padrao = sorted(((nome, aps, outras) for nome, (aps, outras) in por_tipo.items()),
                              key=lambda x: (-x[1], x[0]))

    equips = (db.session.query(TipoEquipamento.nome, db.func.count(Equipamento.id))
              .outerjoin(Equipamento, (Equipamento.tipo_equipamento_id == TipoEquipamento.id) & Equipamento.ativo.is_(True))
              .filter(TipoEquipamento.codigo.is_(None), TipoEquipamento.ativo.is_(True))
              .group_by(TipoEquipamento.nome).order_by(db.func.count(Equipamento.id).desc()).all())
    return salas_sem_padrao, equips


def salvar_pendencias(caminho, pend, log, salas_sem_padrao, equips_fora):
    wb = openpyxl.Workbook()
    cab_fill = PatternFill('solid', fgColor='1A82B8')
    cab_font = Font(bold=True, color='FFFFFF')

    def aba(ws, titulo, colunas, linhas, larguras):
        ws.title = titulo
        ws.append(colunas)
        for c in ws[1]:
            c.fill, c.font = cab_fill, cab_font
            c.alignment = Alignment(vertical='center', wrap_text=True)
        for linha in linhas:
            ws.append(list(linha))
        for i, w in enumerate(larguras):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = 'A2'

    aba(wb.active, 'Pendências da planilha', ['Aba', 'Linha', 'Pendência', 'Detalhe'],
        sorted(pend, key=lambda p: (p[0], p[2], p[1] if isinstance(p[1], int) else 0)), [16, 8, 48, 90])
    aba(wb.create_sheet(), 'Salas sem ambiente padrão',
        ['Tipo de sala no SIGUS', 'Salas em UBS/USF', 'Salas em outras unidades', 'O que fazer'],
        [(n, a, o, 'Reclassificar as salas das UBS em um ambiente padrão' if a else '')
         for n, a, o in salas_sem_padrao], [36, 18, 22, 60])
    aba(wb.create_sheet(), 'Equipamentos fora do catálogo',
        ['Tipo de equipamento no SIGUS', 'Equipamentos cadastrados', 'O que fazer'],
        [(n, q, 'Avaliar inclusão no catálogo de itens') for n, q in equips_fora], [36, 24, 50])
    aba(wb.create_sheet(), 'O que foi importado', ['Ação', 'Detalhe'], log, [30, 100])
    wb.save(caminho)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--planilha', default=PLANILHA_PADRAO)
    parser.add_argument('--aplicar', action='store_true', help='grava no banco (sem isso, só simula)')
    parser.add_argument('--pendencias', help='caminho do .xlsx de pendências para devolver ao Planejamento')
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

    ambientes, itens, kit, unidades, pend = ler_planilha(args.planilha)
    print(f'Planilha: {len(ambientes)} ambientes, {len(itens)} itens, {len(kit)} itens de kit, {len(unidades)} unidades')

    app = create_app()
    with app.app_context():
        print('Banco:', db.engine.url.render_as_string(hide_password=True))
        log = []
        try:
            importar(ambientes, itens, kit, unidades, pend, log)
            salas_sem_padrao, equips_fora = pendencias_sigus()
            resumo = defaultdict(int)
            for acao, _ in log:
                resumo[acao] += 1
            for acao, n in sorted(resumo.items()):
                print(f'  {acao}: {n}')
            for acao, detalhe in log:
                if acao in ('Tipo de sala renomeado', 'Tipo de equipamento casado', 'Kit'):
                    print(f'    {acao}: {detalhe}')
            print(f'  Pendências: {len(pend)}')
            print(f'  Tipos de sala sem padrão em uso: {len(salas_sem_padrao)} '
                  f'({sum(a for _, a, _ in salas_sem_padrao)} salas em UBS/USF)')
            if args.pendencias:
                salvar_pendencias(args.pendencias, pend, log, salas_sem_padrao, equips_fora)
                print('  Pendências salvas em', args.pendencias)
            if args.aplicar:
                db.session.commit()
                print('OK: importação gravada.')
            else:
                db.session.rollback()
                print('Simulação: nada foi gravado (use --aplicar).')
        except Exception:
            db.session.rollback()
            raise


if __name__ == '__main__':
    main()
