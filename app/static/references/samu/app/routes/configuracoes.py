"""Configurações do Sistema SAMU - apenas perfil ADMINISTRADOR."""
import os
import uuid
import mimetypes
from io import BytesIO
from flask import Blueprint, render_template, redirect, url_for, flash, request, abort, jsonify, current_app, send_file
from flask_login import login_required, current_user

from app import db
from app.models.usuario import Usuario, PERFIS, CBOS, ESCOLARIDADES, VINCULOS, TIPOS_VINCULO
from app.models.matricula import MatriculaProfissional
from app.models.unidade_samu import UnidadeSamu
from app.models.tipo_unidade_samu import TipoUnidadeSamu
from app.models.base_descentralizada import BaseDescentralizada
from app.models.veiculo import Veiculo, TIPOS_VEICULO, COMBUSTIVEIS
from app.models.unidade_saude import UnidadeSaude
from app.models.tipo_unidade_saude import TipoUnidadeSaude
from app.models.tipo_ligacao import TipoLigacao
from app.models.origem_ligacao import OrigemLigacao
from app.models.algoritmo_acolhimento import AlgoritmoAcolhimento
from app.models.pergunta_protocolo import PerguntaProtocolo
from app.models.opcao_resposta import OpcaoResposta
from app.models.cid10 import CID10
from app.models.tipo_ocorrencia import TipoOcorrencia
from app.models.motivo_ocorrencia import MotivoOcorrencia
from app.models.intercorrencia import Intercorrencia
from app.models.feriado import Feriado
from app.models.classificacao_risco import ClassificacaoRisco
from app.utils import br_fone

configuracoes_bp = Blueprint('configuracoes', __name__, url_prefix='/configuracoes')


@configuracoes_bp.before_request
@login_required
def _exigir_admin():
    """Só o perfil Administrador acessa Configurações. Os demais levam um chute pro dashboard."""
    if not current_user.pode('configuracoes'):
        flash('Acesso negado. Apenas o perfil Administrador pode acessar as Configurações.', 'warning')
        return redirect(url_for('dashboard.index'))


def _redirect_sucesso(url, msg=None):
    """Redirect após salvar (sem alerta de sucesso)."""
    return redirect(url)


# ══════════════════════════════════════════════════════════
#  ÍNDICE
# ══════════════════════════════════════════════════════════

@configuracoes_bp.route('/')
@login_required
def index():
    """Redireciona para Profissional (primeiro item da Configuração)."""
    return redirect(url_for('configuracoes.profissionais'))


# ══════════════════════════════════════════════════════════
#  PROFISSIONAIS (Usuários com acesso ao sistema)
# ══════════════════════════════════════════════════════════

def _parse_date(val):
    if not val:
        return None
    try:
        from datetime import date
        return date.fromisoformat(val)
    except (ValueError, TypeError):
        return None


@configuracoes_bp.route('/profissionais')
@login_required
def profissionais():
    itens = Usuario.query.order_by(Usuario.nome).all()
    unidades_samu = UnidadeSamu.query.filter_by(ativo=True).order_by(UnidadeSamu.nome).all()
    matriculas_por_usuario = {}
    user_cbo = {}
    for u in itens:
        mats_list = u.matriculas.order_by(MatriculaProfissional.numero).all()
        mats = [{'id': m.id, 'numero': m.numero, 'cbo_label': m.cbo_label,
                 'orgao_emissor': m.orgao_emissor, 'reg_conselho': m.reg_conselho}
                for m in mats_list]
        matriculas_por_usuario[u.id] = mats
        user_cbo[u.id] = (mats_list[0].cbo if mats_list else u.cbo) or ''
    itens = sorted(itens, key=lambda u: (user_cbo.get(u.id, ''), (u.nome or '').lower()))
    return render_template('configuracoes/profissionais_listar.html',
        itens=itens, perfis=PERFIS, cbos=CBOS, unidades_samu=unidades_samu,
        matriculas_por_usuario=matriculas_por_usuario)


def _profissionais_filtrar(itens, nome='', perfil=None, matricula='', cbo=None, ativo=''):
    """Filtra lista de profissionais (mesma lógica do client-side)."""
    if perfil is None:
        perfil = []
    if cbo is None:
        cbo = []
    nome = (nome or '').strip().lower()
    matricula = (matricula or '').strip().lower()
    result = []
    for u in itens:
        if nome and (u.nome or '').lower().find(nome) == -1:
            continue
        if matricula:
            mats_nums = ' '.join(m.numero or '' for m in u.matriculas.all()).lower()
            if matricula not in mats_nums:
                continue
        if ativo and str(int(bool(u.ativo))) != ativo:
            continue
        if perfil:
            ups = set(u.perfis_list)
            if not ups.intersection(perfil):
                continue
        if cbo:
            cbos_u = set()
            for m in u.matriculas.all():
                if m.cbo:
                    cbos_u.add(m.cbo)
            if not cbos_u and u.cbo:
                cbos_u.add(u.cbo)
            if not cbos_u.intersection(cbo):
                continue
        result.append(u)
    return result


@configuracoes_bp.route('/profissionais/export')
@login_required
def profissionais_export():
    """Exporta profissionais para Excel (XLSX) com filtros opcionais."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    itens = Usuario.query.order_by(Usuario.nome).all()
    user_cbo = {}
    for u in itens:
        mats_list = u.matriculas.order_by(MatriculaProfissional.numero).all()
        user_cbo[u.id] = (mats_list[0].cbo if mats_list else u.cbo) or ''
    itens = sorted(itens, key=lambda u: (user_cbo.get(u.id, ''), (u.nome or '').lower()))

    nome = request.args.get('nome', '')
    perfil = [p.strip() for p in request.args.get('perfil', '').split(',') if p.strip()]
    matricula = request.args.get('matricula', '')
    cbo = [c.strip() for c in request.args.get('cbo', '').split(',') if c.strip()]
    ativo = request.args.get('ativo', '')

    itens = _profissionais_filtrar(itens, nome=nome, perfil=perfil, matricula=matricula, cbo=cbo, ativo=ativo)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Profissionais'

    headers = ['Nome', 'E-mail', 'WhatsApp', 'Perfil', 'Usuário', 'Matrícula(s) | CBO', 'Nº Conselho', 'Ativo']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)

    for row_idx, u in enumerate(itens, 2):
        mats = u.matriculas.order_by(MatriculaProfissional.numero).all()
        mat_cbo = '; '.join(
            (m.numero or '') + (' — ' + m.cbo_label if m.cbo else '')
            for m in mats
        ) if mats else (u.cbo_label if u.cbo else '')
        conselho = ''
        if mats and (mats[0].orgao_emissor or mats[0].reg_conselho):
            conselho = ' '.join(filter(None, [mats[0].orgao_emissor, mats[0].reg_conselho]))
        elif u.orgao_emissor or u.reg_conselho:
            conselho = ' '.join(filter(None, [u.orgao_emissor, u.reg_conselho]))
        fone = br_fone(u.whatsapp or u.telefone) if (u.whatsapp or u.telefone) else ''

        ws.cell(row=row_idx, column=1, value=u.nome)
        ws.cell(row=row_idx, column=2, value=u.email or '')
        ws.cell(row=row_idx, column=3, value=fone if fone != '—' else '')
        ws.cell(row=row_idx, column=4, value=u.perfil_label)
        ws.cell(row=row_idx, column=5, value=u.usuario)
        ws.cell(row=row_idx, column=6, value=mat_cbo)
        ws.cell(row=row_idx, column=7, value=conselho)
        ws.cell(row=row_idx, column=8, value='Sim' if u.ativo else 'Não')

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    from datetime import datetime
    fn = f'profissionais_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx'
    return send_file(buf, as_attachment=True, download_name=fn, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@configuracoes_bp.route('/profissionais/novo', methods=['GET', 'POST'], defaults={'id': None})
@configuracoes_bp.route('/profissionais/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def profissional_form(id=None):
    """Apenas Administrador. Para editar o próprio perfil, usar Meu Perfil no menu."""
    item = Usuario.query.get(id) if id else None
    if request.method == 'POST':
        usuario_acesso = request.form.get('usuario', '').strip().lower()
        if not usuario_acesso:
            flash('Usuário de acesso é obrigatório.', 'danger')
            matriculas_err = (MatriculaProfissional.query.filter_by(usuario_id=item.id).order_by(MatriculaProfissional.numero).all() if item else [])
            return render_template('configuracoes/profissional_form.html', item=item,
                perfis=PERFIS, cbos=CBOS, vinculos=VINCULOS, tipos_vinculo=TIPOS_VINCULO, matriculas=matriculas_err)
        conflito = Usuario.query.filter_by(usuario=usuario_acesso).first()
        if conflito and (not item or conflito.id != item.id):
            flash(f'O usuário de acesso "{usuario_acesso}" já está em uso.', 'danger')
            matriculas_err = (MatriculaProfissional.query.filter_by(usuario_id=item.id).order_by(MatriculaProfissional.numero).all() if item else [])
            return render_template('configuracoes/profissional_form.html', item=item,
                perfis=PERFIS, cbos=CBOS, vinculos=VINCULOS, tipos_vinculo=TIPOS_VINCULO, matriculas=matriculas_err)
        if not item:
            item = Usuario(usuario=usuario_acesso)
            item.set_senha('senha')  # Senha padrão
            db.session.add(item)
        else:
            item.usuario = usuario_acesso
        item.nome = request.form.get('nome', '').strip()
        item.apelido = request.form.get('apelido', '').strip() or None
        perfis_sel = request.form.getlist('perfis')
        perfis_sel = [p.strip() for p in perfis_sel if p.strip() and p.strip() in PERFIS]
        if not perfis_sel:
            flash('Marque ao menos um perfil.', 'danger')
            matriculas_err = (MatriculaProfissional.query.filter_by(usuario_id=item.id).order_by(MatriculaProfissional.numero).all() if item else [])
            return render_template('configuracoes/profissional_form.html', item=item,
                perfis=PERFIS, cbos=CBOS, vinculos=VINCULOS, tipos_vinculo=TIPOS_VINCULO, matriculas=matriculas_err)
        item.perfil = ','.join(perfis_sel)
        item.email = request.form.get('email', '').strip() or None
        item.ativo = request.form.get('ativo') == 'on'
        item.cpf = request.form.get('cpf', '').strip() or None
        item.data_nasc = _parse_date(request.form.get('data_nasc'))
        item.nome_mae = request.form.get('nome_mae', '').strip() or None
        whatsapp_raw = request.form.get('whatsapp', '').strip()
        item.whatsapp = ''.join(c for c in whatsapp_raw if c.isdigit()) if whatsapp_raw else None
        nova_senha = request.form.get('nova_senha', '').strip()
        if nova_senha:
            item.set_senha(nova_senha)
        db.session.commit()
        if not id:
            return _redirect_sucesso(url_for('configuracoes.profissional_form', id=item.id), f'Profissional "{item.nome}" salvo!')
        return _redirect_sucesso(url_for('configuracoes.profissionais'), f'Profissional "{item.nome}" salvo!')
    matriculas = []
    if item:
        matriculas = (MatriculaProfissional.query
                      .filter_by(usuario_id=item.id)
                      .order_by(MatriculaProfissional.ativo.desc(), MatriculaProfissional.numero)
                      .all())
    return render_template('configuracoes/profissional_form.html', item=item,
        perfis=PERFIS, cbos=CBOS, vinculos=VINCULOS, tipos_vinculo=TIPOS_VINCULO,
        matriculas=matriculas)


# ─── Foto do Profissional ───────────────────────────────────────────────────
@configuracoes_bp.route('/profissionais/<int:id>/foto', methods=['POST'])
@login_required
def profissional_foto(id):
    usuario = Usuario.query.get_or_404(id)
    foto = request.files.get('foto')
    if foto and foto.filename:
        mime = foto.mimetype or mimetypes.guess_type(foto.filename)[0] or ''
        if not mime.startswith('image/'):
            flash('Apenas imagens são permitidas.', 'danger')
            return redirect(url_for('configuracoes.profissional_form', id=id))
        ext = os.path.splitext(foto.filename)[1].lower() or '.jpg'
        filename = f"{uuid.uuid4().hex}{ext}"
        upload_dir = os.path.join(current_app.static_folder, 'uploads', 'perfis')
        os.makedirs(upload_dir, exist_ok=True)
        foto.save(os.path.join(upload_dir, filename))
        if usuario.foto_perfil:
            try:
                old_path = os.path.join(upload_dir, usuario.foto_perfil)
                if os.path.exists(old_path):
                    os.remove(old_path)
            except OSError:
                pass
        usuario.foto_perfil = filename
        db.session.commit()
        return _redirect_sucesso(url_for('configuracoes.profissional_form', id=id), 'Foto atualizada!')
    return redirect(url_for('configuracoes.profissional_form', id=id))


@configuracoes_bp.route('/profissionais/<int:id>/remover-foto', methods=['POST'])
@login_required
def profissional_remover_foto(id):
    usuario = Usuario.query.get_or_404(id)
    if usuario.foto_perfil:
        upload_dir = os.path.join(current_app.static_folder, 'uploads', 'perfis')
        try:
            old_path = os.path.join(upload_dir, usuario.foto_perfil)
            if os.path.exists(old_path):
                os.remove(old_path)
        except OSError:
            pass
        usuario.foto_perfil = None
        db.session.commit()
        return _redirect_sucesso(url_for('configuracoes.profissional_form', id=id), 'Foto removida.')
    return redirect(url_for('configuracoes.profissional_form', id=id))


# ─── API Matrículas ────────────────────────────────────────────────────────
@configuracoes_bp.route('/profissionais/<int:id>/matriculas', methods=['POST'])
@login_required
def profissional_criar_matricula(id):
    usuario = Usuario.query.get_or_404(id)
    data = request.get_json() or {}
    numero = (data.get('numero') or '').strip()
    if not numero:
        return jsonify(ok=False, erro='Número da matrícula é obrigatório.')
    m = MatriculaProfissional(
        usuario_id=id,
        numero=numero,
        vinculo=data.get('vinculo') or None,
        tipo_vinculo=data.get('tipo_vinculo') or None,
        cbo=data.get('cbo') or None,
        reg_conselho=(data.get('reg_conselho') or '').strip() or None,
        orgao_emissor=(data.get('orgao_emissor') or '').strip() or None,
        ativo=bool(data.get('ativo', True)),
    )
    db.session.add(m)
    db.session.commit()
    return jsonify(ok=True, id=m.id)


@configuracoes_bp.route('/profissionais/<int:uid>/matriculas/<int:mid>', methods=['PUT'])
@login_required
def profissional_atualizar_matricula(uid, mid):
    usuario = Usuario.query.get_or_404(uid)
    m = MatriculaProfissional.query.filter_by(id=mid, usuario_id=uid).first_or_404()
    data = request.get_json() or {}
    numero = (data.get('numero') or '').strip()
    if not numero:
        return jsonify(ok=False, erro='Número da matrícula é obrigatório.')
    m.numero = numero
    m.vinculo = data.get('vinculo') or None
    m.tipo_vinculo = data.get('tipo_vinculo') or None
    m.cbo = data.get('cbo') or None
    m.reg_conselho = (data.get('reg_conselho') or '').strip() or None
    m.orgao_emissor = (data.get('orgao_emissor') or '').strip() or None
    m.ativo = bool(data.get('ativo', True))
    db.session.commit()
    return jsonify(ok=True)


@configuracoes_bp.route('/profissionais/<int:id>/matriculas-json')
@login_required
def profissional_matriculas_json(id):
    """API: retorna matrículas do profissional para o modal Ficha CNES."""
    usuario = Usuario.query.get_or_404(id)
    mats = [{'id': m.id, 'numero': m.numero, 'cbo_label': m.cbo_label,
             'orgao_emissor': m.orgao_emissor, 'reg_conselho': m.reg_conselho}
            for m in usuario.matriculas.order_by(MatriculaProfissional.numero).all()]
    return jsonify(mats)


@configuracoes_bp.route('/profissionais/<int:id>/ficha-cnes')
@login_required
def profissional_ficha_cnes(id):
    """Exibe Ficha CNES do profissional. Aceita: matricula_id, unidade_id, carga_horaria, dt_entrada."""
    usuario = Usuario.query.get_or_404(id)
    matricula_id = request.args.get('matricula_id', type=int)
    if matricula_id and matricula_id > 0:
        mat = MatriculaProfissional.query.filter_by(id=matricula_id, usuario_id=usuario.id).first() or usuario.matricula_principal
    else:
        mat = None  # usa dados diretos do usuário
    unidade_id = request.args.get('unidade_id', type=int)
    unidade = UnidadeSamu.query.get(unidade_id) if unidade_id else usuario.unidade_samu
    carga_horaria = request.args.get('carga_horaria', type=int)
    dt_entrada = _parse_date(request.args.get('dt_entrada', ''))
    return render_template('configuracoes/ficha_cnes.html',
        usuario=usuario, unidade=unidade, matricula=mat,
        carga_horaria=carga_horaria, dt_entrada=dt_entrada)


# ══════════════════════════════════════════════════════════
#  UNIDADES SAMU
# ══════════════════════════════════════════════════════════

def _ordenar_unidades_samu(itens):
    """Ordena: Central primeiro (por apelido); depois viaturas da cidade da Central; depois outras cidades (alfabético); em cada grupo por apelido."""
    central = next((u for u in itens if u.e_central), None)
    central_cidade = (central.cidade or '').strip().lower() if central else ''
    def cidade(u):
        if u.e_central:
            return (u.cidade or '').strip()
        return (u.base.cidade or '').strip() if u.base else ''
    def chave(u):
        apelido_ord = (u.apelido or u.nome or '').strip().lower()
        if u.e_central:
            return (0, '', apelido_ord)
        c = cidade(u)
        c_lower = c.lower()
        if central_cidade and c_lower == central_cidade:
            cidade_key = (0, '')
        else:
            cidade_key = (1, c_lower)
        return (1, cidade_key, apelido_ord)
    return sorted(itens, key=chave)


@configuracoes_bp.route('/unidades-samu')
@login_required
def unidades_samu():
    itens = UnidadeSamu.query.options(db.joinedload(UnidadeSamu.base), db.joinedload(UnidadeSamu.tipo_unidade)).all()
    itens = _ordenar_unidades_samu(itens)
    bases = BaseDescentralizada.query.order_by(BaseDescentralizada.nome).all()
    return render_template('configuracoes/unidades_samu_listar.html',
        itens=itens, titulo='Unidades SAMU', nome_rota='unidades_samu', form_route='unidade_samu_form',
        bases=bases)


@configuracoes_bp.route('/unidades-samu/export')
@login_required
def unidades_samu_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    itens = UnidadeSamu.query.options(db.joinedload(UnidadeSamu.base)).all()
    itens = _ordenar_unidades_samu(itens)
    cidade = request.args.get('cidade', '').strip().lower()
    cnes = ''.join(c for c in (request.args.get('cnes', '') or '') if c.isdigit())
    nome = request.args.get('nome', '').strip().lower()
    apelido = request.args.get('apelido', '').strip().lower()
    tipo = request.args.get('tipo', '')
    base_id = request.args.get('base_id', type=int)
    endereco = request.args.get('endereco', '').strip().lower()
    ativo = request.args.get('ativo', '')
    if cidade or cnes or nome or apelido or tipo or base_id or endereco or ativo:
        result = []
        for i in itens:
            if cidade and cidade not in (i.cidade_uf or '').lower(): continue
            if cnes and cnes not in (i.cnes or ''): continue
            if nome and nome not in (i.nome or '').lower(): continue
            if apelido and apelido not in (i.apelido or '').lower(): continue
            if tipo and (('central' if i.e_central else 'viatura') != tipo): continue
            if base_id and (i.base_id or 0) != base_id: continue
            if endereco and endereco not in (i.endereco_completo or '').lower(): continue
            if ativo and str(int(bool(i.ativo))) != ativo: continue
            result.append(i)
        itens = result
    wb = Workbook()
    ws = wb.active
    ws.title = 'Unidades SAMU'
    headers = ['Cidade/UF', 'CNES', 'Nome', 'Apelido', 'Tipo', 'Base', 'Endereço', 'Telefone', 'Ativo']
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h).font = Font(bold=True)
    for ri, i in enumerate(itens, 2):
        ws.cell(row=ri, column=1, value=i.cidade_uf or '')
        ws.cell(row=ri, column=2, value=i.cnes or '')
        ws.cell(row=ri, column=3, value=i.nome or '')
        ws.cell(row=ri, column=4, value=i.apelido or '')
        ws.cell(row=ri, column=5, value='Central' if i.e_central else 'Viatura')
        ws.cell(row=ri, column=6, value=i.base_nome or '')
        ws.cell(row=ri, column=7, value=i.endereco_completo or '')
        ws.cell(row=ri, column=8, value=br_fone(i.telefone_exibicao) if i.telefone_exibicao else '')
        ws.cell(row=ri, column=9, value='Sim' if i.ativo else 'Não')
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    from datetime import datetime
    return send_file(buf, as_attachment=True, download_name=f'unidades_samu_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@configuracoes_bp.route('/unidades-samu/novo', methods=['GET', 'POST'], defaults={'id': None})
@configuracoes_bp.route('/unidades-samu/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def unidade_samu_form(id=None):
    item = UnidadeSamu.query.options(
        db.joinedload(UnidadeSamu.base),
        db.joinedload(UnidadeSamu.tipo_unidade)
    ).get(id) if id else None
    if request.method == 'POST':
        if not item:
            item = UnidadeSamu()
            db.session.add(item)
        tipo_reg = request.form.get('tipo_registro', 'central') or 'central'
        item.tipo_registro = tipo_reg
        item.cnes = request.form.get('cnes', '').strip() or None
        nome = request.form.get('nome', '').strip()
        if not nome:
            flash('O campo Nome é obrigatório.', 'danger')
            tipos = TipoUnidadeSamu.query.order_by(TipoUnidadeSamu.ativo.desc(), TipoUnidadeSamu.sigla).all()
            bases = BaseDescentralizada.query.order_by(BaseDescentralizada.ativo.desc(), BaseDescentralizada.nome).all()
            return render_template('configuracoes/unidade_samu_form.html', item=item, tipos=tipos, bases=bases)
        item.nome = nome
        item.apelido = request.form.get('apelido', '').strip() or None
        item.ativo = request.form.get('ativo') == 'on'
        if tipo_reg == 'central':
            item.logradouro = request.form.get('logradouro', '').strip() or None
            item.numero = request.form.get('numero', '').strip() or None
            item.complemento = request.form.get('complemento', '').strip() or None
            item.bairro = request.form.get('bairro', '').strip() or None
            item.cidade = request.form.get('cidade', '').strip() or None
            item.uf = request.form.get('uf', '').strip() or None
            item.cep = request.form.get('cep', '').strip() or None
            item.telefone = request.form.get('telefone', '').strip() or None
            item.email = request.form.get('email', '').strip() or None
            item.link_google_maps = request.form.get('link_google_maps', '').strip() or None
            item.base_id = None
            item.tipo_unidade_id = None
        else:
            try:
                base_id = int(request.form.get('base_id') or 0) or None
            except (TypeError, ValueError):
                base_id = None
            try:
                tipo_unidade_id = int(request.form.get('tipo_unidade_id') or 0) or None
            except (TypeError, ValueError):
                tipo_unidade_id = None
            if not base_id or not tipo_unidade_id:
                flash('Para Viatura, é obrigatório selecionar Base e Tipo.', 'danger')
                tipos = TipoUnidadeSamu.query.order_by(TipoUnidadeSamu.ativo.desc(), TipoUnidadeSamu.sigla).all()
                bases = BaseDescentralizada.query.order_by(BaseDescentralizada.ativo.desc(), BaseDescentralizada.nome).all()
                return render_template('configuracoes/unidade_samu_form.html', item=item, tipos=tipos, bases=bases)
            item.base_id = base_id
            item.tipo_unidade_id = tipo_unidade_id
            item.logradouro = item.numero = item.complemento = item.bairro = None
            item.cidade = item.uf = item.cep = item.telefone = item.email = item.link_google_maps = None
        try:
            db.session.commit()
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao salvar: {str(e)}', 'danger')
            tipos = TipoUnidadeSamu.query.order_by(TipoUnidadeSamu.ativo.desc(), TipoUnidadeSamu.sigla).all()
            bases = BaseDescentralizada.query.order_by(BaseDescentralizada.ativo.desc(), BaseDescentralizada.nome).all()
            return render_template('configuracoes/unidade_samu_form.html', item=item, tipos=tipos, bases=bases)
        return _redirect_sucesso(url_for('configuracoes.unidades_samu'), f'Unidade SAMU "{item.nome}" salva!')
    tipos = TipoUnidadeSamu.query.order_by(TipoUnidadeSamu.ativo.desc(), TipoUnidadeSamu.sigla).all()
    bases = BaseDescentralizada.query.order_by(BaseDescentralizada.ativo.desc(), BaseDescentralizada.nome).all()
    if item and item.e_viatura:
        if item.base_id and not any(b.id == item.base_id for b in bases):
            b = BaseDescentralizada.query.get(item.base_id)
            if b:
                bases = [b] + list(bases)
        if item.tipo_unidade_id and not any(t.id == item.tipo_unidade_id for t in tipos):
            t = TipoUnidadeSamu.query.get(item.tipo_unidade_id)
            if t:
                tipos = [t] + list(tipos)
    return render_template('configuracoes/unidade_samu_form.html', item=item, tipos=tipos, bases=bases)


# ══════════════════════════════════════════════════════════
#  VEÍCULOS
# ══════════════════════════════════════════════════════════

@configuracoes_bp.route('/veiculos')
@login_required
def veiculos():
    itens = Veiculo.query.order_by(Veiculo.prefixo).all()
    return render_template('configuracoes/veiculos_listar.html',
        itens=itens, titulo='Veículos', nome_rota='veiculos', form_route='veiculo_form')


@configuracoes_bp.route('/veiculos/export')
@login_required
def veiculos_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    itens = Veiculo.query.order_by(Veiculo.prefixo).all()
    prefixo = request.args.get('prefixo', '').strip().lower()
    marca = request.args.get('marca', '').strip().lower()
    modelo = request.args.get('modelo', '').strip().lower()
    placa = request.args.get('placa', '').strip().lower()
    tipo = request.args.get('tipo', '').strip().lower()
    ativo = request.args.get('ativo', '')
    if prefixo or marca or modelo or placa or tipo or ativo:
        result = [i for i in itens if (not prefixo or prefixo in (i.prefixo or '').lower())
            and (not marca or marca in (i.marca or '').lower()) and (not modelo or modelo in (i.modelo or '').lower())
            and (not placa or placa in (i.placa or '').lower()) and (not tipo or tipo in (i.tipo or '').lower())
            and (not ativo or str(int(bool(i.ativo))) == ativo)]
        itens = result
    wb = Workbook()
    ws = wb.active
    ws.title = 'Veículos'
    for col, h in enumerate(['Prefixo', 'Marca', 'Modelo', 'Placa', 'Tipo', 'Ativo'], 1):
        ws.cell(row=1, column=col, value=h).font = Font(bold=True)
    for ri, i in enumerate(itens, 2):
        ws.cell(row=ri, column=1, value=i.prefixo or '')
        ws.cell(row=ri, column=2, value=i.marca or '')
        ws.cell(row=ri, column=3, value=i.modelo or '')
        ws.cell(row=ri, column=4, value=i.placa or '')
        ws.cell(row=ri, column=5, value=i.tipo or '')
        ws.cell(row=ri, column=6, value='Sim' if i.ativo else 'Não')
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    from datetime import datetime
    return send_file(buf, as_attachment=True, download_name=f'veiculos_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@configuracoes_bp.route('/veiculos/novo', methods=['GET', 'POST'], defaults={'id': None})
@configuracoes_bp.route('/veiculos/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def veiculo_form(id=None):
    item = Veiculo.query.get(id) if id else None
    if request.method == 'POST':
        if not item:
            item = Veiculo()
            db.session.add(item)
        item.prefixo = request.form.get('prefixo', '').strip()
        item.marca = request.form.get('marca', '').strip() or None
        item.modelo = request.form.get('modelo', '').strip() or None
        item.placa = request.form.get('placa', '').strip().upper().replace('-', '') or None
        item.renavam = request.form.get('renavam', '').strip().replace('.', '').replace('-', '') or None
        item.tipo = request.form.get('tipo', '').strip() or None
        item.combustivel = request.form.get('combustivel', '').strip() or None
        item.ano_fabricacao = request.form.get('ano_fabricacao', type=int) or None
        item.ano_modelo = request.form.get('ano_modelo', type=int) or None
        item.ativo = request.form.get('ativo') == 'on'

        db.session.flush()  # garante id para novo registro

        # Upload CRLV (PDF)
        crlv_file = request.files.get('crlv_anexo')
        if crlv_file and crlv_file.filename:
            mime = crlv_file.mimetype or mimetypes.guess_type(crlv_file.filename)[0] or ''
            if mime == 'application/pdf' or (crlv_file.filename or '').lower().endswith('.pdf'):
                upload_dir = os.path.join(current_app.static_folder, 'uploads', 'veiculos')
                os.makedirs(upload_dir, exist_ok=True)
                filename = f"crlv_{item.id}.pdf"
                crlv_file.save(os.path.join(upload_dir, filename))
                if item.crlv_anexo and item.crlv_anexo != filename:
                    try:
                        old = os.path.join(upload_dir, item.crlv_anexo)
                        if os.path.exists(old):
                            os.remove(old)
                    except OSError:
                        pass
                item.crlv_anexo = filename

        db.session.commit()
        return _redirect_sucesso(url_for('configuracoes.veiculos'), f'Veículo "{item.prefixo}" salvo!')
    return render_template('configuracoes/veiculo_form.html', item=item,
        tipos_veiculo=TIPOS_VEICULO, combustiveis=COMBUSTIVEIS)


@configuracoes_bp.route('/veiculos/<int:id>/crlv')
@login_required
def veiculo_crlv(id):
    """Download do CRLV anexado."""
    item = Veiculo.query.get_or_404(id)
    if not item.crlv_anexo:
        abort(404)
    from flask import send_from_directory
    return send_from_directory(
        os.path.join(current_app.static_folder, 'uploads', 'veiculos'),
        item.crlv_anexo, as_attachment=True, download_name=f'CRLV_{item.prefixo.replace(" ", "_")}.pdf'
    )


# ══════════════════════════════════════════════════════════
#  TIPOS DE UNIDADE SAMU
# ══════════════════════════════════════════════════════════

@configuracoes_bp.route('/tipos-unidade-samu')
@login_required
def tipos_unidade_samu():
    itens = TipoUnidadeSamu.query.order_by(TipoUnidadeSamu.sigla).all()
    return render_template('configuracoes/tipos_unidade_samu_listar.html',
        itens=itens, titulo='Tipos de Unidades SAMU')


@configuracoes_bp.route('/tipos-unidade-samu/export')
@login_required
def tipos_unidade_samu_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    itens = TipoUnidadeSamu.query.order_by(TipoUnidadeSamu.sigla).all()
    nome = request.args.get('nome', '').strip().lower()
    sigla = request.args.get('sigla', '').strip().lower()
    ativo = request.args.get('ativo', '')
    if nome or sigla or ativo:
        result = [i for i in itens if (not nome or nome in (i.nome or '').lower()) and (not sigla or sigla in (i.sigla or '').lower()) and (not ativo or str(int(bool(i.ativo))) == ativo)]
        itens = result
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tipos Unidade SAMU'
    for col, h in enumerate(['Nome', 'Sigla', 'Ativo'], 1):
        ws.cell(row=1, column=col, value=h).font = Font(bold=True)
    for ri, i in enumerate(itens, 2):
        ws.cell(row=ri, column=1, value=i.nome or '')
        ws.cell(row=ri, column=2, value=i.sigla or '')
        ws.cell(row=ri, column=3, value='Sim' if i.ativo else 'Não')
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    from datetime import datetime
    return send_file(buf, as_attachment=True, download_name=f'tipos_unidade_samu_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@configuracoes_bp.route('/tipos-unidade-samu/novo', methods=['GET', 'POST'], defaults={'id': None})
@configuracoes_bp.route('/tipos-unidade-samu/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def tipo_unidade_samu_form(id=None):
    item = TipoUnidadeSamu.query.get(id) if id else None
    if request.method == 'POST':
        if not item:
            item = TipoUnidadeSamu()
            db.session.add(item)
        item.sigla = request.form.get('sigla', '').strip()
        if not item.predefinido:
            item.nome = request.form.get('nome', '').strip()
        item.icon = request.form.get('icon', '').strip() or None
        try:
            item.ordem_prioridade = int(request.form.get('ordem_prioridade') or 0)
        except (TypeError, ValueError):
            item.ordem_prioridade = 0
        item.ativo = request.form.get('ativo') == 'on'
        db.session.commit()
        return _redirect_sucesso(url_for('configuracoes.tipos_unidade_samu'), 'Tipo de unidade salvo!')
    return render_template('configuracoes/tipo_unidade_samu_form.html', item=item)


# ══════════════════════════════════════════════════════════
#  BASES DESCENTRALIZADAS
# ══════════════════════════════════════════════════════════

@configuracoes_bp.route('/bases-descentralizadas')
@login_required
def bases_descentralizadas():
    itens = BaseDescentralizada.query.order_by(BaseDescentralizada.cidade, BaseDescentralizada.nome).all()
    return render_template('configuracoes/bases_descentralizadas_listar.html',
        itens=itens, titulo='Bases', nome_rota='bases_descentralizadas', form_route='base_descentralizada_form')


@configuracoes_bp.route('/bases-descentralizadas/export')
@login_required
def bases_descentralizadas_export():
    from openpyxl import Workbook
    itens = BaseDescentralizada.query.order_by(BaseDescentralizada.cidade, BaseDescentralizada.nome).all()
    cidade = request.args.get('cidade', '').strip().lower()
    nome = request.args.get('nome', '').strip().lower()
    endereco = request.args.get('endereco', '').strip().lower()
    email = request.args.get('email', '').strip().lower()
    ativo = request.args.get('ativo', '')
    if cidade or nome or endereco or email or ativo:
        result = []
        for i in itens:
            if cidade and cidade not in ((i.cidade or '') + (i.uf or '')).lower(): continue
            if nome and nome not in (i.nome or '').lower(): continue
            if endereco and endereco not in (i.endereco_completo or '').lower(): continue
            if email and email not in (i.email or '').lower(): continue
            if ativo and str(int(bool(i.ativo))) != ativo: continue
            result.append(i)
        itens = result
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = 'Bases'
    for col, h in enumerate(['Cidade/UF', 'Nome', 'Endereço', 'Telefone', 'E-mail', 'Ativo'], 1):
        ws.cell(row=1, column=col, value=h).font = Font(bold=True)
    for ri, i in enumerate(itens, 2):
        cid_uf = (i.cidade or '') + '/' + (i.uf or '') if (i.cidade or i.uf) else ''
        ws.cell(row=ri, column=1, value=cid_uf.strip('/'))
        ws.cell(row=ri, column=2, value=i.nome or '')
        ws.cell(row=ri, column=3, value=i.endereco_completo or '')
        ws.cell(row=ri, column=4, value=br_fone(i.telefone) if i.telefone else '')
        ws.cell(row=ri, column=5, value=i.email or '')
        ws.cell(row=ri, column=6, value='Sim' if i.ativo else 'Não')
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    from datetime import datetime
    return send_file(buf, as_attachment=True, download_name=f'bases_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@configuracoes_bp.route('/bases-descentralizadas/novo', methods=['GET', 'POST'], defaults={'id': None})
@configuracoes_bp.route('/bases-descentralizadas/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def base_descentralizada_form(id=None):
    item = BaseDescentralizada.query.get(id) if id else None
    if request.method == 'POST':
        if not item:
            item = BaseDescentralizada()
            db.session.add(item)
        item.nome = request.form.get('nome', '').strip()
        item.logradouro = request.form.get('logradouro', '').strip() or None
        item.numero = request.form.get('numero', '').strip() or None
        item.complemento = request.form.get('complemento', '').strip() or None
        item.bairro = request.form.get('bairro', '').strip() or None
        item.cidade = request.form.get('cidade', '').strip() or None
        item.uf = request.form.get('uf', '').strip() or None
        item.cep = request.form.get('cep', '').strip() or None
        item.telefone = request.form.get('telefone', '').strip() or None
        item.email = request.form.get('email', '').strip() or None
        item.link_google_maps = request.form.get('link_google_maps', '').strip() or None
        item.ativo = request.form.get('ativo') == 'on'
        db.session.commit()
        return _redirect_sucesso(url_for('configuracoes.bases_descentralizadas'), f'Base "{item.nome}" salva!')
    return render_template('configuracoes/base_descentralizada_form.html', item=item)


# ══════════════════════════════════════════════════════════
#  TIPOS DE UNIDADES DE SAÚDE
# ══════════════════════════════════════════════════════════

@configuracoes_bp.route('/tipos-unidade-saude')
@login_required
def tipos_unidade_saude():
    itens = TipoUnidadeSaude.query.order_by(TipoUnidadeSaude.sigla).all()
    return render_template('configuracoes/tipos_unidade_saude_listar.html', itens=itens, titulo='Tipos de Unidades de Saúde')


@configuracoes_bp.route('/tipos-unidade-saude/export')
@login_required
def tipos_unidade_saude_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    itens = TipoUnidadeSaude.query.order_by(TipoUnidadeSaude.sigla).all()
    sigla = request.args.get('sigla', '').strip().lower()
    nome = request.args.get('nome', '').strip().lower()
    ativo = request.args.get('ativo', '')
    if sigla or nome or ativo:
        result = [i for i in itens if (not sigla or sigla in (i.sigla or '').lower()) and (not nome or nome in (i.nome or '').lower()) and (not ativo or str(int(bool(i.ativo))) == ativo)]
        itens = result
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tipos Unidade Saúde'
    for col, h in enumerate(['Sigla', 'Nome', 'Ativo'], 1):
        ws.cell(row=1, column=col, value=h).font = Font(bold=True)
    for ri, i in enumerate(itens, 2):
        ws.cell(row=ri, column=1, value=i.sigla or '')
        ws.cell(row=ri, column=2, value=i.nome or '')
        ws.cell(row=ri, column=3, value='Sim' if i.ativo else 'Não')
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    from datetime import datetime
    return send_file(buf, as_attachment=True, download_name=f'tipos_unidade_saude_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@configuracoes_bp.route('/tipos-unidade-saude/novo', methods=['GET', 'POST'], defaults={'id': None})
@configuracoes_bp.route('/tipos-unidade-saude/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def tipo_unidade_saude_form(id=None):
    item = TipoUnidadeSaude.query.get(id) if id else None
    if request.method == 'POST':
        if not item:
            item = TipoUnidadeSaude()
            db.session.add(item)
        item.sigla = request.form.get('sigla', '').strip()
        item.nome = request.form.get('nome', '').strip()
        item.ativo = request.form.get('ativo') == 'on'
        db.session.commit()
        return _redirect_sucesso(url_for('configuracoes.tipos_unidade_saude'))
    return render_template('configuracoes/tipo_unidade_saude_form.html', item=item)


# ══════════════════════════════════════════════════════════
#  UNIDADES DE SAÚDE
# ══════════════════════════════════════════════════════════

def _ordenar_unidades_saude(itens):
    """Ordena por cidade da Central primeiro, depois alfabético por cidade e nome."""
    central = UnidadeSamu.query.filter_by(tipo_registro='central', ativo=True).first()
    central_cidade = (central.cidade or '').strip().lower() if central else ''
    def chave(u):
        c = (u.cidade or u.municipio or '').strip().lower()
        if central_cidade and c == central_cidade:
            cidade_key = (0, '')
        else:
            cidade_key = (1, c)
        return (cidade_key, (u.nome or '').strip().lower())
    return sorted(itens, key=chave)


@configuracoes_bp.route('/unidades-saude')
@login_required
def unidades_saude():
    itens = UnidadeSaude.query.options(db.joinedload(UnidadeSaude.tipo_unidade)).all()
    itens = _ordenar_unidades_saude(itens)
    tipos = TipoUnidadeSaude.query.order_by(TipoUnidadeSaude.sigla).all()
    return render_template('configuracoes/unidades_saude_listar.html',
        itens=itens, titulo='Unidades de Saúde', nome_rota='unidades_saude', form_route='unidade_saude_form',
        tipos_unidade_saude=tipos)


@configuracoes_bp.route('/unidades-saude/export')
@login_required
def unidades_saude_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    itens = UnidadeSaude.query.options(db.joinedload(UnidadeSaude.tipo_unidade)).all()
    itens = _ordenar_unidades_saude(itens)
    cidade = request.args.get('cidade', '').strip().lower()
    cnes = ''.join(c for c in (request.args.get('cnes', '') or '') if c.isdigit())
    nome = request.args.get('nome', '').strip().lower()
    tipo_ids = [int(x) for x in (request.args.get('tipo_id', '') or '').split(',') if x.strip().isdigit()]
    endereco = request.args.get('endereco', '').strip().lower()
    ativo = request.args.get('ativo', '')
    if cidade or cnes or nome or tipo_ids or endereco or ativo:
        result = []
        for i in itens:
            if cidade and cidade not in (i.cidade_uf or '').lower(): continue
            if cnes and cnes not in (i.cnes or ''): continue
            if nome and nome not in (i.nome or '').lower(): continue
            if tipo_ids and (i.tipo_unidade_id or 0) not in tipo_ids: continue
            if endereco and endereco not in (i.endereco_completo or '').lower(): continue
            if ativo and str(int(bool(i.ativo))) != ativo: continue
            result.append(i)
        itens = result
    wb = Workbook()
    ws = wb.active
    ws.title = 'Unidades Saúde'
    headers = ['Cidade/UF', 'CNES', 'Nome', 'Tipo', 'Horário', 'Telefone', 'Endereço', 'Ativo']
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h).font = Font(bold=True)
    for ri, i in enumerate(itens, 2):
        ws.cell(row=ri, column=1, value=i.cidade_uf or '')
        ws.cell(row=ri, column=2, value=i.cnes or '')
        ws.cell(row=ri, column=3, value=i.nome or '')
        ws.cell(row=ri, column=4, value=i.tipo_unidade.sigla if i.tipo_unidade else '')
        ws.cell(row=ri, column=5, value=i.hora_funcionamento or '')
        ws.cell(row=ri, column=6, value=br_fone(i.telefone) if i.telefone else '')
        ws.cell(row=ri, column=7, value=i.endereco_completo or '')
        ws.cell(row=ri, column=8, value='Sim' if i.ativo else 'Não')
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    from datetime import datetime
    return send_file(buf, as_attachment=True, download_name=f'unidades_saude_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _parse_time(s):
    if not s or not s.strip():
        return None
    try:
        from datetime import datetime
        return datetime.strptime(s.strip()[:5], '%H:%M').time()
    except (ValueError, TypeError):
        return None


@configuracoes_bp.route('/unidades-saude/novo', methods=['GET', 'POST'], defaults={'id': None})
@configuracoes_bp.route('/unidades-saude/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def unidade_saude_form(id=None):
    item = UnidadeSaude.query.get(id) if id else None
    if request.method == 'POST':
        if not item:
            item = UnidadeSaude()
            db.session.add(item)
        item.cnes = request.form.get('cnes', '').strip() or None
        item.nome = request.form.get('nome', '').strip()
        item.tipo_unidade_id = request.form.get('tipo_unidade_id', type=int) or None
        item.logradouro = request.form.get('logradouro', '').strip() or None
        item.numero = request.form.get('numero', '').strip() or None
        item.complemento = request.form.get('complemento', '').strip() or None
        item.bairro = request.form.get('bairro', '').strip() or None
        item.cidade = request.form.get('cidade', '').strip() or None
        item.uf = request.form.get('uf', '').strip() or None
        item.cep = request.form.get('cep', '').strip() or None
        item.telefone = request.form.get('telefone', '').strip() or None
        item.email = request.form.get('email', '').strip() or None
        item.link_google_maps = request.form.get('link_google_maps', '').strip() or None
        item.funcionamento_24h = request.form.get('funcionamento_24h') == 'on'
        item.hora_inicio = _parse_time(request.form.get('hora_inicio')) if not item.funcionamento_24h else None
        item.hora_fim = _parse_time(request.form.get('hora_fim')) if not item.funcionamento_24h else None
        item.ativo = request.form.get('ativo') == 'on'
        db.session.commit()
        return _redirect_sucesso(url_for('configuracoes.unidades_saude'))
    tipos = TipoUnidadeSaude.query.order_by(TipoUnidadeSaude.sigla).all()
    return render_template('configuracoes/unidade_saude_form.html', item=item, tipos=tipos)


# ══════════════════════════════════════════════════════════
#  HELPERS GENÉRICOS PARA CRUD SIMPLES
# ══════════════════════════════════════════════════════════

def _crud_simples(Model, nome_rota, titulo, campo_nome='nome', campos_extra=None):
    """Factory para rotas CRUD de tabelas simples (nome, descricao, ativo)."""
    @configuracoes_bp.route(f'/{nome_rota}')
    @login_required
    def listar():
        itens = Model.query.order_by(getattr(Model, campo_nome)).all()
        cols = [(campo_nome, 'Nome')]
        if campos_extra:
            cols.extend(campos_extra)
        return render_template('configuracoes/crud_listar.html',
            itens=itens, titulo=titulo, nome_rota=nome_rota, colunas=cols)

    @configuracoes_bp.route(f'/{nome_rota}/novo', methods=['GET', 'POST'])
    @configuracoes_bp.route(f'/{nome_rota}/<int:id>/editar', methods=['GET', 'POST'])
    @login_required
    def form(id=None):
        item = Model.query.get(id) if id else None
        if request.method == 'POST':
            if not item:
                item = Model()
                db.session.add(item)
            setattr(item, campo_nome, request.form.get(campo_nome, '').strip())
            if hasattr(item, 'descricao'):
                item.descricao = request.form.get('descricao', '').strip() or None
            if hasattr(item, 'ativo'):
                item.ativo = request.form.get('ativo') == 'on'
            if campos_extra:
                for attr, _ in campos_extra:
                    if hasattr(item, attr):
                        setattr(item, attr, request.form.get(attr, '').strip() or None)
            db.session.commit()
            return _redirect_sucesso(url_for(f'configuracoes.{listar.__name__}'), f'{titulo[:-1]} salvo!')
        return render_template('configuracoes/crud_form.html',
            item=item, titulo=titulo, nome_rota=nome_rota, campo_nome=campo_nome, campos_extra=campos_extra or [])

    return listar, form


# Registrar rotas para entidades simples
_CONF_SIMPLES = [
    (TipoLigacao, 'tipos-ligacao', 'Tipos de Ligação'),
    (OrigemLigacao, 'origens-ligacao', 'Origens de Ligação'),
    (AlgoritmoAcolhimento, 'algoritmos-acolhimento', 'Algoritmos de Acolhimento', 'codigo', [('nome', 'Nome'), ('cor', 'Cor')]),
    (CID10, 'cid10', 'CID-10', 'codigo', [('descricao', 'Descrição')]),
    (TipoOcorrencia, 'tipos-ocorrencia', 'Tipos de Ocorrência'),
    (MotivoOcorrencia, 'motivos-ocorrencia', 'Motivos de Ocorrência'),
    (Intercorrencia, 'intercorrencias', 'Intercorrências'),
]

# Algoritmo e CID10 têm estrutura diferente - codigo/descricao
# Vou criar rotas manuais para cada um para ter mais controle
def _register_simple_cruds():
    # Tipo Ligação
    @configuracoes_bp.route('/tipos-ligacao')
    @login_required
    def tipos_ligacao():
        itens = TipoLigacao.query.order_by(TipoLigacao.nome).all()
        return render_template('configuracoes/tipos_ligacao_listar.html', itens=itens, titulo='Tipos de Ligação')

    @configuracoes_bp.route('/tipos-ligacao/export')
    @login_required
    def tipos_ligacao_export():
        from openpyxl import Workbook
        from openpyxl.styles import Font
        itens = TipoLigacao.query.order_by(TipoLigacao.nome).all()
        nome = request.args.get('nome', '').strip().lower()
        descricao = request.args.get('descricao', '').strip().lower()
        ativo = request.args.get('ativo', '')
        if nome or descricao or ativo:
            result = [i for i in itens if (not nome or nome in (i.nome or '').lower()) and (not descricao or descricao in (i.descricao or '').lower()) and (not ativo or str(int(bool(i.ativo))) == ativo)]
            itens = result
        wb = Workbook()
        ws = wb.active
        ws.title = 'Tipos Ligação'
        for col, h in enumerate(['Nome', 'Descrição', 'Ativo'], 1):
            ws.cell(row=1, column=col, value=h).font = Font(bold=True)
        for ri, i in enumerate(itens, 2):
            ws.cell(row=ri, column=1, value=i.nome or '')
            ws.cell(row=ri, column=2, value=i.descricao or '')
            ws.cell(row=ri, column=3, value='Sim' if i.ativo else 'Não')
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        from datetime import datetime
        return send_file(buf, as_attachment=True, download_name=f'tipos_ligacao_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @configuracoes_bp.route('/tipos-ligacao/novo', methods=['GET', 'POST'], defaults={'id': None})
    @configuracoes_bp.route('/tipos-ligacao/<int:id>/editar', methods=['GET', 'POST'])
    @login_required
    def tipo_ligacao_form(id=None):
        item = TipoLigacao.query.get(id) if id else None
        if request.method == 'POST':
            if not item:
                item = TipoLigacao()
                db.session.add(item)
            item.nome = request.form.get('nome', '').strip()
            item.descricao = request.form.get('descricao', '').strip() or None
            item.permite_encerrar_direto = request.form.get('permite_encerrar_direto') == 'on'
            item.eh_trote = request.form.get('eh_trote') == 'on'
            item.ativo = request.form.get('ativo') == 'on'
            db.session.commit()
            return _redirect_sucesso(url_for('configuracoes.tipos_ligacao'), 'Tipo de ligação salvo!')
        return render_template('configuracoes/tipo_ligacao_form.html', item=item)

    # Origem Ligação
    @configuracoes_bp.route('/origens-ligacao')
    @login_required
    def origens_ligacao():
        itens = OrigemLigacao.query.order_by(OrigemLigacao.nome).all()
        return render_template('configuracoes/origens_ligacao_listar.html', itens=itens, titulo='Origens de Ligação')

    @configuracoes_bp.route('/origens-ligacao/export')
    @login_required
    def origens_ligacao_export():
        from openpyxl import Workbook
        from openpyxl.styles import Font
        itens = OrigemLigacao.query.order_by(OrigemLigacao.nome).all()
        nome = request.args.get('nome', '').strip().lower()
        descricao = request.args.get('descricao', '').strip().lower()
        ativo = request.args.get('ativo', '')
        if nome or descricao or ativo:
            result = [i for i in itens if (not nome or nome in (i.nome or '').lower()) and (not descricao or descricao in (i.descricao or '').lower()) and (not ativo or str(int(bool(i.ativo))) == ativo)]
            itens = result
        wb = Workbook()
        ws = wb.active
        ws.title = 'Origens Ligação'
        for col, h in enumerate(['Nome', 'Descrição', 'Ativo'], 1):
            ws.cell(row=1, column=col, value=h).font = Font(bold=True)
        for ri, i in enumerate(itens, 2):
            ws.cell(row=ri, column=1, value=i.nome or '')
            ws.cell(row=ri, column=2, value=i.descricao or '')
            ws.cell(row=ri, column=3, value='Sim' if i.ativo else 'Não')
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        from datetime import datetime
        return send_file(buf, as_attachment=True, download_name=f'origens_ligacao_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @configuracoes_bp.route('/origens-ligacao/novo', methods=['GET', 'POST'], defaults={'id': None})
    @configuracoes_bp.route('/origens-ligacao/<int:id>/editar', methods=['GET', 'POST'])
    @login_required
    def origem_ligacao_form(id=None):
        item = OrigemLigacao.query.get(id) if id else None
        if request.method == 'POST':
            if not item:
                item = OrigemLigacao()
                db.session.add(item)
            item.nome = request.form.get('nome', '').strip()
            item.descricao = request.form.get('descricao', '').strip() or None
            item.preenche_endereco_unidade_saude = request.form.get('preenche_endereco_unidade_saude') == 'on'
            item.ativo = request.form.get('ativo') == 'on'
            db.session.commit()
            return _redirect_sucesso(url_for('configuracoes.origens_ligacao'), 'Origem de ligação salva!')
        return render_template('configuracoes/origem_ligacao_form.html', item=item)

    # API: buscar protocolos por nome (para Tom Select na tela TARM)
    @configuracoes_bp.route('/api/protocolos-buscar')
    @login_required
    def api_protocolos_buscar():
        """Retorna protocolos cujo nome contém o texto informado (q=)."""
        q = (request.args.get('q') or '').strip()
        if not q:
            protos = AlgoritmoAcolhimento.query.filter_by(ativo=True).order_by(AlgoritmoAcolhimento.nome).limit(50).all()
        else:
            from sqlalchemy import func
            protos = AlgoritmoAcolhimento.query.filter(
                AlgoritmoAcolhimento.ativo == True,
                func.lower(AlgoritmoAcolhimento.nome).contains(q.lower())
            ).order_by(AlgoritmoAcolhimento.nome).limit(50).all()
        return jsonify([{'id': p.id, 'nome': p.nome} for p in protos])

    @configuracoes_bp.route('/api/protocolo/<int:id>')
    @login_required
    def api_protocolo_detalhe(id):
        """Retorna um protocolo com suas perguntas e opções (para preenchimento guiado)."""
        p = AlgoritmoAcolhimento.query.filter_by(id=id, ativo=True).first()
        if not p:
            return jsonify({'error': 'Protocolo não encontrado'}), 404
        perguntas = []
        for prg in sorted(p.perguntas.all(), key=lambda x: x.ordem):
            if prg.tipo in (PerguntaProtocolo.TIPO_MULTIPLA, PerguntaProtocolo.TIPO_SELECTBOX):
                opt = [{'texto': o.texto, 'pontuacao': getattr(o, 'pontuacao', 0) or 0} for o in sorted(prg.opcoes.all(), key=lambda x: x.ordem)]
            else:
                opt = None
            perguntas.append({'id': prg.id, 'secao': prg.secao, 'enunciado': prg.enunciado, 'tipo': prg.tipo, 'obrigatoria': prg.obrigatoria, 'opcoes': opt})
        return jsonify({'id': p.id, 'nome': p.nome, 'codigo': p.codigo, 'perguntas': perguntas, 'graus_risco': p.graus_risco or []})

    # Algoritmo Acolhimento (protocolo com palavras-chave e perguntas)
    @configuracoes_bp.route('/algoritmos-acolhimento')
    @login_required
    def algoritmos_acolhimento():
        itens = AlgoritmoAcolhimento.query.order_by(AlgoritmoAcolhimento.nome).all()
        return render_template('configuracoes/algoritmos_acolhimento_listar.html',
            itens=itens, titulo='Protocolos de Acolhimento')

    @configuracoes_bp.route('/algoritmos-acolhimento/novo', methods=['GET', 'POST'], defaults={'id': None})
    @configuracoes_bp.route('/algoritmos-acolhimento/<int:id>/editar', methods=['GET', 'POST'])
    @login_required
    def algoritmo_acolhimento_form(id=None):
        item = AlgoritmoAcolhimento.query.get(id) if id else None
        if request.method == 'POST':
            if not item:
                item = AlgoritmoAcolhimento()
                db.session.add(item)
            nome = request.form.get('nome', '').strip()
            item.nome = nome
            item.descricao = request.form.get('descricao', '').strip() or None
            item.ativo = request.form.get('ativo') == 'on'
            # Graus de risco: grau_N_min, grau_N_max, grau_N_grau
            graus = []
            gidx = 0
            while True:
                gmin = request.form.get(f'grau_{gidx}_min', '').strip()
                gmax = request.form.get(f'grau_{gidx}_max', '').strip()
                ggrau = request.form.get(f'grau_{gidx}_grau', '').strip()
                if not ggrau and not gmin and not gmax:
                    break
                try:
                    graus.append({
                        'min': int(gmin) if gmin else 0,
                        'max': int(gmax) if gmax else 0,
                        'grau': ggrau or '—'
                    })
                except ValueError:
                    pass
                gidx += 1
            item.graus_risco = graus if graus else None
            if not item.codigo and nome:
                import re
                slug = re.sub(r'[^a-zA-Z0-9]+', '_', nome).strip('_').upper()[:20]
                item.codigo = slug or 'P'
            db.session.flush()
            # Perguntas: pergunta_N_*
            perguntas_atuais = list(item.perguntas.order_by(PerguntaProtocolo.ordem).all())
            idx = 0
            while True:
                enc = request.form.get(f'pergunta_{idx}_enunciado', '').strip()
                if not enc:
                    break
                secao = request.form.get(f'pergunta_{idx}_secao', '').strip() or None
                tipo = request.form.get(f'pergunta_{idx}_tipo', PerguntaProtocolo.TIPO_SELECTBOX) or PerguntaProtocolo.TIPO_SELECTBOX
                obrigatoria = request.form.get(f'pergunta_{idx}_obrigatoria') == 'on'
                if idx < len(perguntas_atuais):
                    prg = perguntas_atuais[idx]
                    prg.secao = secao
                    prg.enunciado = enc[:10000]  # TEXT permite enunciados longos (ex.: descrição do procedimento)
                    prg.tipo = tipo
                    prg.obrigatoria = obrigatoria
                    prg.ordem = idx
                    for op in prg.opcoes.all():
                        db.session.delete(op)
                else:
                    prg = PerguntaProtocolo(algoritmo_id=item.id, secao=secao, enunciado=enc[:10000], tipo=tipo, obrigatoria=obrigatoria, ordem=idx)
                    db.session.add(prg)
                    db.session.flush()
                if tipo in (PerguntaProtocolo.TIPO_MULTIPLA, PerguntaProtocolo.TIPO_SELECTBOX):
                    opcoes_linhas = (request.form.get(f'pergunta_{idx}_opcoes') or '').strip().split('\n')
                    for oi, opt in enumerate(opcoes_linhas):
                        opt = opt.strip()
                        if not opt:
                            continue
                        # Formato: "texto" ou "texto|pontuação"
                        partes = opt.split('|', 1)
                        texto = partes[0].strip()[:200] if partes[0] else ''
                        pontuacao = 0
                        if len(partes) > 1 and partes[1].strip():
                            try:
                                pontuacao = int(partes[1].strip())
                            except ValueError:
                                pass
                        if texto:
                            opc = OpcaoResposta(pergunta_id=prg.id, texto=texto, pontuacao=pontuacao, ordem=oi)
                            db.session.add(opc)
                idx += 1
            for prg in perguntas_atuais[idx:]:
                db.session.delete(prg)
            db.session.commit()
            return _redirect_sucesso(url_for('configuracoes.algoritmos_acolhimento'), 'Protocolo salvo!')
        perguntas = list(item.perguntas.order_by(PerguntaProtocolo.ordem).all()) if item else []
        graus_risco = (item.graus_risco or []) if item else []
        return render_template('configuracoes/algoritmo_form.html', item=item, TIPOS=PerguntaProtocolo.TIPOS,
            perguntas=perguntas, graus_risco=graus_risco)

    @configuracoes_bp.route('/algoritmos-acolhimento/<int:id>/fluxo', methods=['GET', 'POST'])
    @login_required
    def algoritmo_fluxo(id):
        item = AlgoritmoAcolhimento.query.get_or_404(id)
        if request.method == 'GET':
            return redirect(url_for('configuracoes.algoritmo_acolhimento_form', id=id, tab='fluxo'))
        if request.method == 'POST':
            perguntas_ordem = {p.id: p.ordem for p in item.perguntas.order_by(PerguntaProtocolo.ordem).all()}
            for opc in OpcaoResposta.query.filter(OpcaoResposta.pergunta_id.in_(perguntas_ordem.keys())).all():
                val = request.form.get(f'opcao_{opc.id}_proxima', type=int)
                ordem_atual = perguntas_ordem.get(opc.pergunta_id, -1)
                ordem_destino = perguntas_ordem.get(val) if val else None
                if val and (ordem_destino is None or ordem_destino <= ordem_atual):
                    val = None  # rejeitar: mesma pergunta, anterior ou ID inválido
                opc.proxima_pergunta_id = val if val else None
            db.session.commit()
            flash('Fluxo salvo!', 'success')
            return redirect(url_for('configuracoes.algoritmo_acolhimento_form', id=id, tab='fluxo'))
        perguntas = list(item.perguntas.order_by(PerguntaProtocolo.ordem).all())
        return render_template('configuracoes/algoritmo_fluxo.html', item=item, perguntas=perguntas)

    # CID-10 (codigo, descricao) — busca sob demanda (não carrega 14k linhas)
    @configuracoes_bp.route('/cid10')
    @login_required
    def cid10():
        tem_registros = CID10.query.limit(1).first() is not None
        return render_template('configuracoes/cid10_listar.html',
            tem_registros=tem_registros, titulo='CID-10')

    @configuracoes_bp.route('/cid10/api/buscar')
    @login_required
    def cid10_api_buscar():
        from sqlalchemy import or_
        busca = request.args.get('q', '').strip()[:100]
        ativo = request.args.get('ativo', '')
        lim = int(request.args.get('limit', 300))
        limit = min(lim, 20000) if not busca else min(lim, 500)
        q = CID10.query.order_by(CID10.codigo)
        if ativo:
            q = q.filter(CID10.ativo == (ativo == '1'))
        if busca:
            term = f'%{busca}%'
            q = q.filter(or_(
                CID10.codigo.ilike(term),
                CID10.descricao.ilike(term)
            ))
        itens = q.limit(limit).all()
        return jsonify([{
            'id': i.id,
            'codigo': i.codigo or '',
            'descricao': i.descricao or '',
            'ativo': i.ativo,
        } for i in itens])

    @configuracoes_bp.route('/cid10/importar', methods=['POST'])
    @login_required
    def cid10_importar():
        from app.cid10_import import importar_todos
        try:
            r = importar_todos(current_app)
        except Exception as e:
            flash(f'Erro na importação: {e}', 'danger')
            return redirect(url_for('configuracoes.cid10'))
        if r['erros']:
            for e in r['erros']:
                flash(e, 'danger')
        c10 = r['cid10']
        co = r['cido']
        msgs = []
        if c10['novos'] or c10['atualizados']:
            msgs.append(f"CID-10: {c10['novos']} novos, {c10['atualizados']} atualizados")
        if co['novos'] or co['atualizados']:
            msgs.append(f"CID-O: {co['novos']} novos, {co['atualizados']} atualizados")
        if msgs:
            flash('Importação concluída: ' + ' | '.join(msgs), 'success')
        elif not r['erros']:
            flash('Nenhum dado novo para importar.', 'info')
        return redirect(url_for('configuracoes.cid10'))

    @configuracoes_bp.route('/cid10/export')
    @login_required
    def cid10_export():
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from sqlalchemy import or_
        busca = request.args.get('busca', '').strip()[:100]
        ativo = request.args.get('ativo', '')
        q = CID10.query.order_by(CID10.codigo)
        if ativo:
            q = q.filter(CID10.ativo == (ativo == '1'))
        if busca:
            term = f'%{busca}%'
            q = q.filter(or_(CID10.codigo.ilike(term), CID10.descricao.ilike(term)))
        itens = q.all()
        wb = Workbook()
        ws = wb.active
        ws.title = 'CID-10'
        for col, h in enumerate(['Código', 'Descrição', 'Ativo'], 1):
            ws.cell(row=1, column=col, value=h).font = Font(bold=True)
        for ri, i in enumerate(itens, 2):
            ws.cell(row=ri, column=1, value=i.codigo or '')
            ws.cell(row=ri, column=2, value=i.descricao or '')
            ws.cell(row=ri, column=3, value='Sim' if i.ativo else 'Não')
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        from datetime import datetime
        return send_file(buf, as_attachment=True, download_name=f'cid10_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @configuracoes_bp.route('/cid10/novo', methods=['GET', 'POST'], defaults={'id': None})
    @configuracoes_bp.route('/cid10/<int:id>/editar', methods=['GET', 'POST'])
    @login_required
    def cid10_form(id=None):
        item = CID10.query.get(id) if id else None
        if request.method == 'POST':
            if not item:
                item = CID10()
                db.session.add(item)
            from app.cid10_import import _limpar_codigo
            item.codigo = _limpar_codigo(request.form.get('codigo', ''))
            item.descricao = (request.form.get('descricao', '') or '').strip()
            item.ativo = request.form.get('ativo') == 'on'
            db.session.commit()
            return _redirect_sucesso(url_for('configuracoes.cid10'), 'CID-10 salvo!')
        return render_template('configuracoes/cid10_form.html', item=item)

    # Tipo Ocorrência
    @configuracoes_bp.route('/tipos-ocorrencia')
    @login_required
    def tipos_ocorrencia():
        itens = TipoOcorrencia.query.order_by(TipoOcorrencia.nome).all()
        return render_template('configuracoes/tipos_ocorrencia_listar.html',
            itens=itens, titulo='Tipos de Ocorrência', form_route='tipo_ocorrencia_form')

    @configuracoes_bp.route('/tipos-ocorrencia/novo', methods=['GET', 'POST'], defaults={'id': None})
    @configuracoes_bp.route('/tipos-ocorrencia/<int:id>/editar', methods=['GET', 'POST'])
    @login_required
    def tipo_ocorrencia_form(id=None):
        item = TipoOcorrencia.query.get(id) if id else None
        if request.method == 'POST':
            if not item:
                item = TipoOcorrencia()
                db.session.add(item)
            item.nome = request.form.get('nome', '').strip()
            item.descricao = request.form.get('descricao', '').strip() or None
            item.ativo = request.form.get('ativo') == 'on'
            db.session.commit()
            return _redirect_sucesso(url_for('configuracoes.tipos_ocorrencia'), 'Tipo de ocorrência salvo!')
        return render_template('configuracoes/crud_form.html', item=item, titulo='Tipo de Ocorrência', nome_rota='tipos_ocorrencia', campo_nome='nome')

    # Motivo Ocorrência
    @configuracoes_bp.route('/motivos-ocorrencia')
    @login_required
    def motivos_ocorrencia():
        itens = MotivoOcorrencia.query.outerjoin(MotivoOcorrencia.tipo_ocorrencia).order_by(TipoOcorrencia.nome, MotivoOcorrencia.nome).all()
        tipos = TipoOcorrencia.query.filter_by(ativo=True).order_by(TipoOcorrencia.nome).all()
        return render_template('configuracoes/motivos_ocorrencia_listar.html',
            itens=itens, titulo='Motivos de Ocorrência', form_route='motivo_ocorrencia_form', tipos=tipos)

    @configuracoes_bp.route('/motivos-ocorrencia/novo', methods=['GET', 'POST'], defaults={'id': None})
    @configuracoes_bp.route('/motivos-ocorrencia/<int:id>/editar', methods=['GET', 'POST'])
    @login_required
    def motivo_ocorrencia_form(id=None):
        item = MotivoOcorrencia.query.get(id) if id else None
        tipos = TipoOcorrencia.query.filter_by(ativo=True).all()
        if request.method == 'POST':
            if not item:
                item = MotivoOcorrencia()
                db.session.add(item)
            item.nome = request.form.get('nome', '').strip()
            item.descricao = request.form.get('descricao', '').strip() or None
            item.tipo_ocorrencia_id = request.form.get('tipo_ocorrencia_id', type=int) or None
            item.ativo = request.form.get('ativo') == 'on'
            db.session.commit()
            return _redirect_sucesso(url_for('configuracoes.motivos_ocorrencia'), 'Motivo de ocorrência salvo!')
        return render_template('configuracoes/motivo_ocorrencia_form.html', item=item, tipos=tipos)

    # Classificação de Risco
    @configuracoes_bp.route('/classificacoes-risco')
    @login_required
    def classificacoes_risco():
        itens = ClassificacaoRisco.query.order_by(ClassificacaoRisco.ordem, ClassificacaoRisco.nome).all()
        return render_template('configuracoes/classificacoes_risco_listar.html',
            itens=itens, titulo='Classificação de Risco', form_route='classificacao_risco_form')

    @configuracoes_bp.route('/classificacoes-risco/novo', methods=['GET', 'POST'], defaults={'id': None})
    @configuracoes_bp.route('/classificacoes-risco/<int:id>/editar', methods=['GET', 'POST'])
    @login_required
    def classificacao_risco_form(id=None):
        item = ClassificacaoRisco.query.get(id) if id else None
        if request.method == 'POST':
            if not item:
                item = ClassificacaoRisco()
                db.session.add(item)
            item.nome = request.form.get('nome', '').strip()
            item.cor = (request.form.get('cor') or '#6c757d').strip()[:20]
            try:
                item.ordem = int(request.form.get('ordem') or 0)
            except (TypeError, ValueError):
                item.ordem = 0
            item.ativo = request.form.get('ativo') == 'on'
            db.session.commit()
            return _redirect_sucesso(url_for('configuracoes.classificacoes_risco'), 'Classificação de risco salva!')
        return render_template('configuracoes/classificacao_risco_form.html', item=item)

    # Intercorrência
    @configuracoes_bp.route('/intercorrencias')
    @login_required
    def intercorrencias():
        itens = Intercorrencia.query.order_by(Intercorrencia.nome).all()
        return render_template('configuracoes/intercorrencias_listar.html',
            itens=itens, titulo='Intercorrências', form_route='intercorrencia_form')

    @configuracoes_bp.route('/intercorrencias/novo', methods=['GET', 'POST'], defaults={'id': None})
    @configuracoes_bp.route('/intercorrencias/<int:id>/editar', methods=['GET', 'POST'])
    @login_required
    def intercorrencia_form(id=None):
        item = Intercorrencia.query.get(id) if id else None
        if request.method == 'POST':
            if not item:
                item = Intercorrencia()
                db.session.add(item)
            item.nome = request.form.get('nome', '').strip()
            item.descricao = request.form.get('descricao', '').strip() or None
            item.ativo = request.form.get('ativo') == 'on'
            item.disponivel_cancelamento = request.form.get('disponivel_cancelamento') == 'on'
            db.session.commit()
            return _redirect_sucesso(url_for('configuracoes.intercorrencias'), 'Intercorrência salva!')
        return render_template('configuracoes/intercorrencia_form.html', item=item)

    # Feriados
    @configuracoes_bp.route('/feriados')
    @login_required
    def feriados():
        from app.models.feriado import TIPOS_FOLGA, NATUREZAS
        itens = Feriado.query.order_by(Feriado.data).all()
        return render_template('configuracoes/feriados_listar.html',
            itens=itens, titulo='Feriados', tipos_folga=TIPOS_FOLGA, naturezas=NATUREZAS)

    @configuracoes_bp.route('/feriados/novo', methods=['GET', 'POST'], defaults={'id': None})
    @configuracoes_bp.route('/feriados/<int:id>/editar', methods=['GET', 'POST'])
    @login_required
    def feriado_form(id=None):
        from app.models.feriado import TIPOS_FOLGA, NATUREZAS
        item = Feriado.query.get(id) if id else None
        if request.method == 'POST':
            if not item:
                item = Feriado()
            item.data = _parse_date(request.form.get('data'))
            item.nome = request.form.get('nome', '').strip()
            item.tipo_folga = request.form.get('tipo_folga', 'feriado') or 'feriado'
            if item.tipo_folga == 'data_comemorativa':
                item.natureza = 'nacional'  # placeholder, não exibido na UI
                item.emoji = request.form.get('emoji', '').strip()[:20] or None
            else:
                item.natureza = request.form.get('natureza', 'nacional') or 'nacional'
                item.emoji = None
            if item.tipo_folga == 'data_comemorativa':
                item.numero_decreto = None
                item.link_decreto = None
                item.cor_primaria = (request.form.get('cor_primaria') or '').strip()[:7] or None
                item.cor_secundaria = (request.form.get('cor_secundaria') or '').strip()[:7] or None
                item.cor_fonte_primaria = (request.form.get('cor_fonte_primaria') or '').strip()[:7] or None
                item.cor_fonte_secundaria = (request.form.get('cor_fonte_secundaria') or '').strip()[:7] or None
            else:
                item.cor_primaria = item.cor_secundaria = item.cor_fonte_primaria = item.cor_fonte_secundaria = None
                item.emoji = None
                item.numero_decreto = request.form.get('numero_decreto', '').strip() or None
                item.link_decreto = request.form.get('link_decreto', '').strip() or None
            item.descricao = request.form.get('descricao', '').strip() or None
            item.ativo = request.form.get('ativo') == 'on'
            if not item.data:
                flash('Data é obrigatória.', 'danger')
                return render_template('configuracoes/feriado_form.html', item=item, tipos_folga=TIPOS_FOLGA, naturezas=NATUREZAS)
            if not item.nome:
                flash('Evento é obrigatório.', 'danger')
                return render_template('configuracoes/feriado_form.html', item=item, tipos_folga=TIPOS_FOLGA, naturezas=NATUREZAS)
            outro = Feriado.query.filter(Feriado.data == item.data)
            if item.id:
                outro = outro.filter(Feriado.id != item.id)
            if outro.first():
                flash(f'Já existe um feriado cadastrado para a data {item.data.strftime("%d/%m/%Y")}.', 'danger')
                return render_template('configuracoes/feriado_form.html', item=item, tipos_folga=TIPOS_FOLGA, naturezas=NATUREZAS)
            if not item.id:
                db.session.add(item)
            db.session.commit()
            if request.form.get('salvar_e_novo'):
                return _redirect_sucesso(url_for('configuracoes.feriado_form'), f'Feriado "{item.nome}" salvo!')
            return _redirect_sucesso(url_for('configuracoes.feriados'), f'Feriado "{item.nome}" salvo!')
        return render_template('configuracoes/feriado_form.html', item=item, tipos_folga=TIPOS_FOLGA, naturezas=NATUREZAS)

_register_simple_cruds()
